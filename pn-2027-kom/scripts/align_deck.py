"""Alinha o KOM (v3 do usuário, com slides WBS 14–17) ao WBS v3: datas reprogramadas, lotes, nome padronizado, EDT."""
import sys, re, copy
from datetime import datetime
import openpyxl
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.oxml.ns import qn
DECK, XLSX, OUT = sys.argv[1:4]
NAME_UP = 'AQUISIÇÃO E INSTALAÇÃO DE RL, RT E TRIPSAVER'
W = openpyxl.load_workbook(XLSX, data_only=True)['WBS_Cronograma']
def fd(x): return x.strftime('%d/%m/%Y') if isinstance(x, datetime) else str(x)
# ---------- leitura do WBS
rows = {}
for r in range(5, W.max_row + 1):
    rows[r] = dict(code=W.cell(r, 1).value, name=(W.cell(r, 2).value or ''), dur=W.cell(r, 5).value, ini=W.cell(r, 6).value, fim=W.cell(r, 7).value)
sections = {}      # '3'/'4'/'5' -> row
blocks = {}        # n -> dict(head_row, tasks[10 rows])  (último sub-bloco de cada número)
for r, d in rows.items():
    if str(d['code']) in ('3', '4', '5') and d['name'].startswith('REGIONAL'): sections[str(d['code'])] = r
    m = re.match(r'BLOCO (\d+) ', d['name'])
    if m: blocks[int(m.group(1))] = dict(head=r, tasks=list(range(r + 1, r + 11)))
top = list(range(5, 22))

prs = Presentation(DECK)
S = list(prs.slides)
def set_runs(par, text):
    rs = par.runs
    if not rs: rs = [par.add_run()]
    rs[0].text = text
    for x in rs[1:]: x._r.getparent().remove(x._r)
def cell_set(cell, text): set_runs(cell.text_frame.paragraphs[0], text)
def tables(sl): return sorted([s for s in sl.shapes if s.has_table], key=lambda s: s.left)
def dur(v): return str(int(v)) if isinstance(v, (int, float)) else str(v)

def fill_block_rows(tbl, start_row, n):
    """preenche (dur, início, término) do bloco n a partir de start_row (linha 'BLOCO n') e as 10 tarefas."""
    b = blocks[n]
    h = rows[b['head']]
    cs = tbl.rows[start_row].cells
    cell_set(cs[1], dur(h['dur'])); cell_set(cs[2], fd(h['ini'])); cell_set(cs[3], fd(h['fim']))
    for k, tr in enumerate(b['tasks'], 1):
        t = rows[tr]; cs = tbl.rows[start_row + k].cells
        assert cs[0].text.strip() == t['name'].strip(), (cs[0].text, t['name'])
        cell_set(cs[1], dur(t['dur'])); cell_set(cs[2], fd(t['ini'])); cell_set(cs[3], fd(t['fim']))
def fill_section(tbl, ri, key):
    t = rows[sections[key]]; cs = tbl.rows[ri].cells
    cell_set(cs[1], dur(t['dur'])); cell_set(cs[2], fd(t['ini'])); cell_set(cs[3], fd(t['fim']))

# ---------- S14: gerenciamento + planejamento (esquerda)
tl, tr_ = tables(S[13])
tb = tl.table
# remove linhas excedentes (a WBS v3 tem 17 linhas de topo; o slide tinha 19)
need = len(top)
while len(tb.rows) - 1 > need:
    tb._tbl.remove(tb._tbl.tr_lst[-1])
for i, r in enumerate(top, 1):
    d = rows[r]; cs = tb.rows[i].cells
    is_sum = isinstance(d['code'], int) or str(d['code']).isdigit()
    label = f"{d['code']} · {d['name'].strip()}" if is_sum else f"   {d['code']} {d['name'].strip()}"
    cell_set(cs[0], label); cell_set(cs[1], dur(d['dur'])); cell_set(cs[2], fd(d['ini'])); cell_set(cs[3], fd(d['fim']))
tl.height = Emu(sum(r.height for r in tb.rows))
# direita: Norte (B1, B2)
tb = tr_.table
fill_section(tb, 1, '3'); fill_block_rows(tb, 2, 1); fill_block_rows(tb, 13, 2)
# S15: B3 | Sul (B12) ; B13
tl, tr_ = tables(S[14])
fill_block_rows(tl.table, 1, 3); fill_section(tl.table, 12, '5'); fill_block_rows(tl.table, 13, 12)
fill_block_rows(tr_.table, 1, 13)
# S16: Centro (B5, B4) ; B6, B7
tl, tr_ = tables(S[15])
fill_section(tl.table, 1, '4'); fill_block_rows(tl.table, 2, 5); fill_block_rows(tl.table, 13, 4)
fill_block_rows(tr_.table, 1, 6); fill_block_rows(tr_.table, 12, 7)
# S17: B8, B9 ; B10, B11
tl, tr_ = tables(S[16])
fill_block_rows(tl.table, 1, 8); fill_block_rows(tl.table, 12, 9)
fill_block_rows(tr_.table, 1, 10); fill_block_rows(tr_.table, 12, 11)

# ---------- S10: EDT do xlsx e GAT em todos os blocos
edt = {1: '3.1', 2: '3.2', 3: '3.3', 4: '4.1 / 4.2', 5: '4.3', 6: '4.4', 7: '4.5', 8: '4.6', 9: '4.7', 10: '4.8', 11: '4.9', 12: '5.1 / 5.2', 13: '5.3'}
t10 = [s for s in S[9].shapes if s.has_table][0]
cell_set(t10.table.rows[0].cells[0], 'EDT')
for n in range(1, 14): cell_set(t10.table.rows[n].cells[0], edt[n])
t10.table.columns[0].width = Emu(t10.table.columns[0].width + 160000)
t10.table.columns[1].width = Emu(t10.table.columns[1].width - 160000)
side = [s for s in S[9].shapes if s.name == 'TextBox 12'][0]
paras = side.text_frame.paragraphs
for p_ in paras:
    if not p_.runs: continue
    tx = p_.text
    tx = re.sub(r'^1\.n\.(\d+)', r'EDT.\1', tx)
    tx = tx.replace('Chamado / Atend. GAT – só RL (8 d)', 'Chamado / Atend. GAT (8 d)')
    set_runs(p_, tx)
root = [s for s in S[9].shapes if s.name == 'TextBox 9'][0]
ps = root.text_frame._txBody.findall(qn('a:p'))
if ps[0].xpath('string(.)').strip() == '1': root.text_frame._txBody.remove(ps[0])

# ---------- cabeçalho do grupo com 7 blocos (slide 11) — cabia em 2 linhas
for sl in S[10:13]:
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith('BLOCOS 1, 2, 5'):
            set_runs(sh.text_frame.paragraphs[0], 'BLOCOS 1, 2, 5, 6, 7, 8 e 12  ·  SIGCO 92118, 89267, 90811 (RT), 92095, 92101, 70838 e 91979  —  execução de 30 dias')
# ---------- eyebrow, rodapé, numeração
ref_eb = [s for s in S[8].shapes if s.name == 'Eyebrow'][0]
for sl in S[13:17]:
    eb = [s for s in sl.shapes if s.name == 'Eyebrow'][0]
    txb = eb.text_frame._txBody
    for p_ in txb.findall(qn('a:p')): txb.remove(p_)
    txb.append(copy.deepcopy(ref_eb.text_frame._txBody.find(qn('a:p'))))
    eb.width = ref_eb.width
S = list(prs.slides)
for idx, sl in enumerate(S, 1):
    for f in [s for s in sl.shapes if s.name == 'Rodapé']:
        set_runs(f.text_frame.paragraphs[0], f'{idx:02d}  ·  {NAME_UP}  ·  PN 2027')
prs.save(OUT)
print('saved', OUT, len(S))
