"""WBS v3: adiciona Gerenciamento (PMBOK) e Planejamento/Suprimentos (lotes) no topo, igual ao deck; corrige 'Parte III' -> 'Parte II'."""
import sys, copy, re, datetime as dt
import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
SRC, OUT = sys.argv[1:3]
SHIFT = 7
wb = openpyxl.load_workbook(SRC)
w = wb['WBS_Cronograma']
last = w.max_row
# estilos de referência (antes de mover)
def grab(r): return [copy.copy(w.cell(r, c)._style) for c in range(1, 12)]
st_sum, st_item = grab(15), grab(6)
h5, h6 = w.row_dimensions[5].height, w.row_dimensions[6].height
# 1) move o corpo (a partir da linha 15) 7 linhas para baixo; fórmulas internas são traduzidas
w.move_range(f'A15:K{last}', rows=SHIFT, translate=True)
# 2) outras abas: referências à WBS com linha >= 15 sobem 7
pat = re.compile(r"(WBS_Cronograma!\$?[A-Z]{1,2}\$?)(\d+)(:\$?[A-Z]{1,2}\$?)?(\d+)?")
def bump(m):
    r1 = int(m.group(2)); n1 = r1 + SHIFT if r1 >= 15 else r1
    if m.group(3):
        r2 = int(m.group(4)); n2 = r2 + SHIFT if r2 >= 15 else r2
        return f'{m.group(1)}{n1}{m.group(3)}{n2}'
    return f'{m.group(1)}{n1}'
for n in ('Curva_Orcamentaria', 'Resumo_Blocos'):
    ws = wb[n]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and 'WBS_Cronograma!' in c.value: c.value = pat.sub(bump, c.value)
# 3) limpa as linhas 5..21 e escreve o topo novo
for r in range(5, 22):
    for c in range(1, 12):
        cell = w.cell(r, c); cell.value = None; cell.comment = None
D = lambda y, m, d: dt.datetime(y, m, d)
# localiza linhas de início de execução / EC dos blocos (para fórmulas dos lotes e do CM)
ec_rows, blk = [], {}
exec_rows = {}
cur = None
for r in range(22, w.max_row + 1):
    b = w.cell(r, 2).value
    if isinstance(b, str):
        m = re.match(r'BLOCO (\d+) ', b)
        if m: cur = int(m.group(1)); blk.setdefault(cur, []).append(r)
        if b.strip() == 'Encerramento Contábil': ec_rows.append(r)
        if b.strip() == 'Executar Obras': exec_rows.setdefault(cur, []).append(r)
        if b.strip() == 'Equipamentos em Operação': pass
        if re.match(r'BLOCO 5 ', b): w.cell(r, 2).value = b.replace('(Parte III)', '(Parte II)')
first_exec = lambda ns: '{' if False else 'MIN(' + ','.join(f'F{exec_rows[n][0]}' for n in ns) + ')'
rows = [
 # (r, code, name, dur, start, end, resp, status, kind)
 (5, 1, 'GERENCIAMENTO DO PROJETO (PMBOK)', '=G5-F5', '=MIN(F6:F10)', '=MAX(G6:G10)', None, None, 'sum'),
 (6, '1.1', 'Transfer Meeting (TM)', 0, D(2026, 10, 6), '=F6+E6', 'ASGP / ASPO', 'Concluído', 'item'),
 (7, '1.2', 'Kickoff Meeting (KOM)', 0, D(2026, 10, 29), '=F7+E7', 'GP / PMO', 'A iniciar', 'item'),
 (8, '1.3', 'Project Review Meeting (PRM) / linha de base', 0, D(2026, 12, 2), '=F8+E8', 'GP / PMO', 'A iniciar', 'item'),
 (9, '1.4', 'Follow-up Meetings (FUM) mensais', 'mensal', D(2027, 1, 15), D(2027, 8, 16), 'GP', 'A iniciar', 'item'),
 (10, '1.5', 'Close Meeting (CM)', 0, '=MAX(' + ','.join(f'G{r}' for r in ec_rows) + ')+7', '=F10+E10', 'GP / PMO', 'A iniciar', 'item'),
 (11, 2, 'PLANEJAMENTO E SUPRIMENTOS', '=G11-F11', '=MIN(F12:F21)', '=MAX(G12:G21)', None, None, 'sum'),
 (12, '2.1', 'Criar PMA de materiais específicos (RL, RT, TripSaver)', 30, D(2026, 11, 3), '=F12+E12', 'GP / Suprimentos', 'A iniciar', 'item'),
 (13, '2.2', 'Realizar estudo Site Survey', 30, D(2026, 11, 9), '=F13+E13', 'DCMD / Projetista', 'A iniciar', 'item'),
 (14, '2.3', 'Elaborar projetos executivos', 60, D(2026, 11, 3), '=F14+E14', 'Projetista', 'A iniciar', 'item'),
 (15, '2.4', 'Analisar projetos executivos (COPC)', 15, D(2027, 1, 4), '=F15+E15', 'COPC', 'A iniciar', 'item'),
 (16, '2.5', 'Check de consistência orçamento x projeto executivo', 10, D(2027, 1, 20), '=F16+E16', 'GP', 'A iniciar', 'item'),
 (17, '2.6', 'Report da análise para ASPO/DCMD', 5, D(2027, 2, 1), '=F17+E17', 'GP', 'A iniciar', 'item'),
 (18, '2.7', 'Receber equipamentos - Lote 1 (B1, B2 e B4 a B12)', 0, D(2027, 4, 5), '=F18+E18', 'DEOP / Suprimentos', 'A iniciar', 'item'),
 (19, '2.8', 'Receber equipamentos - Lote 2 (B3 e B13)', 0, '=' + first_exec([3, 13]) + '-7', '=F19+E19', 'DEOP / Suprimentos', 'A iniciar', 'item'),
 (20, '2.9', 'Emitir O.S. e mobilizar EPS - Lote 1', 5, '=G20-E20', '=' + first_exec([1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12]) + '-1', 'DCMD / Contratação', 'A iniciar', 'item'),
 (21, '2.10', 'Emitir O.S. e mobilizar EPS - Lote 2', 5, '=G21-E21', '=' + first_exec([3, 13]) + '-1', 'DCMD / Contratação', 'A iniciar', 'item'),
]
for r, code, name, dur, s_, e_, resp, status, kind in rows:
    st = st_sum if kind == 'sum' else st_item
    for c in range(1, 12): w.cell(r, c)._style = copy.copy(st[c - 1])
    w.cell(r, 1).value = str(code)
    w.cell(r, 2).value = ('' if kind == 'sum' else '   ') + name
    w.cell(r, 5).value = dur; w.cell(r, 6).value = s_; w.cell(r, 7).value = e_
    w.cell(r, 8).value = resp; w.cell(r, 11).value = status
    w.row_dimensions[r].height = h5 if kind == 'sum' else h6
# datas de entrada em azul já vêm do estilo 'item' (verifica as de fórmula)
black = copy.copy(w['F17'].font)
for r in (10, 19, 20, 21):
    w.cell(r, 6).font = Font(name=w.cell(r, 6).font.name, size=w.cell(r, 6).font.size, color='000000')
w['F18'].comment = Comment('Equipamentos (RL, RT e TripSaver) disponíveis somente após 05/04/2027: nenhuma execução inicia antes de 06/04/2027.', 'Claude')
w['F19'].comment = Comment('Lote 2 (B3 e B13): recebimento 7 dias antes do início da execução (blocos de junho).', 'Claude')
# Premissa 5
pf = wb['Premissas_Fontes']
for rr in range(1, pf.max_row + 1):
    v = pf.cell(rr, 1).value
    if isinstance(v, str) and v.startswith('5. Equipamentos'):
        pf.cell(rr, 1).value = ('5. Equipamentos (RL/RT/TripSaver) disponíveis somente após 05/04/2027: nenhuma execução de bloco inicia antes de 06/04/2027 '
                                '(Lote 1 recebido em 05/04/2027; Lote 2 – B3 e B13 – recebido 7 dias antes do início); O.S. da EPS emitida nos 5 dias que antecedem o início de cada lote.')
    if isinstance(v, str) and v.startswith('1. Cronograma considera'):
        pf.cell(rr, 1).value = v + ' Gerenciamento (PMBOK) e planejamento/suprimentos constam da WBS apenas como marcos e atividades de apoio.'
# formatação condicional (folga negativa) refeita
w.conditional_formatting = type(w.conditional_formatting)()
red = Font(color='C00000', bold=True)
w.conditional_formatting.add(f'J22:J{w.max_row}', CellIsRule(operator='lessThan', formula=['0'], font=red))
wb.save(OUT)
print('ok', w.max_row, ec_rows[:3], exec_rows)
