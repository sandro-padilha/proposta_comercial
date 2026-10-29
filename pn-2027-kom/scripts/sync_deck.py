"""Sincroniza o KOM (v2) com o cronograma reprogramado (equipamentos disponíveis só após 05/04/2027)."""
import sys, copy, calendar, re
from datetime import date, datetime, timedelta
import openpyxl
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

DECK, XLSX, OUT = sys.argv[1:4]
NAME_UP = 'AQUISIÇÃO E INSTALAÇÃO DE RL, RT E TRIPSAVER'
NAVY, ORANGE, TEAL, TEAL2, DARKB, LIME, GREY, GREY2, GREEN, TXT = '1F1F4C', 'F47B20', '009FC2', '66C5DB', '006199', 'B5E622', '989898', 'C9C9C9', '628C0E', '4D4D4D'
RED = 'C00000'

# ------------------------------------------------------------------ dados do xlsx (valores calculados)
wv = openpyxl.load_workbook(XLSX, data_only=True)
W = wv['WBS_Cronograma']; RB = wv['Resumo_Blocos']; CV = wv['Curva_Orcamentaria']
def D(x): return x.date() if isinstance(x, datetime) else x
blocks = {}      # n -> dict (último sub-bloco com aquele número)
r = 1
while r <= W.max_row:
    nm = W.cell(r, 2).value
    m = re.match(r'BLOCO (\d+) ', nm) if isinstance(nm, str) else None
    if m:
        n = int(m.group(1))
        rows = {}
        keys = ['exec', 'conc', 'fisc', 'cad', 'tci', 'gat', 'com', 'op', 'et', 'ec']
        for k, rr in zip(keys, range(r + 1, r + 11)):
            rows[k] = (D(W.cell(rr, 6).value), D(W.cell(rr, 7).value), W.cell(rr, 5).value)
        op_row = r + 8
        blocks[n] = dict(n=n, sigco=W.cell(r, 4).value, name=nm, rows=rows, meta=D(W.cell(op_row, 9).value), folga=W.cell(op_row, 10).value)
        r += 11
    else:
        r += 1
assert sorted(blocks) == list(range(1, 14)), sorted(blocks)
resumo = {}
for rr in range(4, 17):
    resumo[RB.cell(rr, 1).value] = dict(sigco=RB.cell(rr, 2).value, reg=RB.cell(rr, 3).value, tipo=RB.cell(rr, 4).value, esc=RB.cell(rr, 5).value)
curva = []
for rr in range(5, 17):
    curva.append(dict(sigco=CV.cell(rr, 1).value, nome=CV.cell(rr, 2).value, reg=CV.cell(rr, 3).value, total=CV.cell(rr, 4).value,
                      m=[CV.cell(rr, c).value or 0 for c in range(9, 21)]))
tot_m = [CV.cell(17, c).value or 0 for c in range(9, 21)]
reg_m = {CV.cell(rr, 2).value: [CV.cell(rr, c).value or 0 for c in range(9, 21)] + [CV.cell(rr, 21).value] for rr in (22, 23, 24)}
tot = CV.cell(17, 21).value
acc_pct = [CV.cell(19, c).value for c in range(9, 21)]
LIM = date(2027, 4, 5)
for n, b in blocks.items(): assert b['rows']['exec'][0] > LIM, n

def brl(v, dec=0):
    s = f'{v:,.{dec}f}'; return s.replace(',', 'X').replace('.', ',').replace('X', '.')
def dm(d): return d.strftime('%d/%m')
def dmy(d): return d.strftime('%d/%m/%Y')
def eom(d): return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])
late = [n for n, b in blocks.items() if b['folga'] < 0]
first_start = min(b['rows']['exec'][0] for b in blocks.values())
last_op = max(b['rows']['op'][0] for b in blocks.values())

prs = Presentation(DECK)
S = list(prs.slides)

# ------------------------------------------------------------------ helpers
def shp(sl, name, nth=0): return [s for s in sl.shapes if s.name == name][nth]
def set_runs(par, text):
    rs = par.runs; rs[0].text = text
    for x in rs[1:]: x._r.getparent().remove(x._r)
def set_paras(sh, texts, size=None):
    tb = sh.text_frame._txBody; ps = tb.findall(qn('a:p'))
    while len(ps) < len(texts):
        n_ = copy.deepcopy(ps[-1]); ps[-1].addnext(n_); ps = tb.findall(qn('a:p'))
    for i, p_ in enumerate(ps):
        if i < len(texts): set_runs(sh.text_frame.paragraphs[i], texts[i])
        else: tb.remove(p_)
    if size:
        for pa in sh.text_frame.paragraphs:
            for x in pa.runs: x.font.size = Pt(size)
def cell_set(cell, text, color=None, size=None, tight=False):
    p_ = cell.text_frame.paragraphs[0]
    if p_.runs: set_runs(p_, text)
    else:
        rr = p_.add_run(); rr.text = text
    rn = p_.runs[0]
    if color: rn.font.color.rgb = rgb(color)
    if size: rn.font.size = Pt(size)
    if tight:
        tcp = cell._tc.find(qn('a:tcPr')); tcp.set('marL', '9144'); tcp.set('marR', '9144')
def drop(sh): sh._element.getparent().remove(sh._element)
def rgb(h): return RGBColor.from_string(h)

def rect(sl, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    return s
def text(sl, x, y, w, h, t, size=8, bold=False, color=TXT, align='l', anchor='m'):
    tb = sl.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(36000); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE}[anchor]
    p_ = tf.paragraphs[0]; p_.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
    rr = p_.add_run(); rr.text = t; rr.font.size = Pt(size); rr.font.bold = bold; rr.font.name = 'DM Sans'; rr.font.color.rgb = rgb(color)
    return tb

def clear_body(sl, keep=('Title 2', 'Eyebrow', 'Rodapé')):
    for s in list(sl.shapes):
        if s.name not in keep: drop(s)

def month_axis(sl, x0, x1, y, h, m_first, m_last, y_strip_h, label_h=None, size=9):
    """cabeçalho de meses (navy) + faixas alternadas; devolve X(date)."""
    d0 = date(2027, m_first, 1); d1 = date(2027, m_last + 1, 1) if m_last < 12 else date(2028, 1, 1)
    days = (d1 - d0).days
    X = lambda d: x0 + (x1 - x0) * (d - d0).days / days
    mn = {1: 'Jan', 2: 'Fev', 3: 'Mar', 4: 'Abr', 5: 'Mai', 6: 'Jun', 7: 'Jul', 8: 'Ago', 9: 'Set', 10: 'Out', 11: 'Nov', 12: 'Dez'}
    for i, m in enumerate(range(m_first, m_last + 1)):
        a = date(2027, m, 1); b = date(2027, m + 1, 1) if m < 12 else date(2028, 1, 1)
        if i % 2 == 0: rect(sl, X(a), y + h, X(b) - X(a), y_strip_h, 'F7F8FB')
        rect(sl, X(a), y, X(b) - X(a), h, NAVY)
        text(sl, X(a), y, X(b) - X(a), h, f'{mn[m]}/27', size, True, 'FFFFFF', 'c')
    return X

# ------------------------------------------------------------------ slide 9 — cronograma macro
macro = S[8]
clear_body(macro)
lab = {1: ('RL ×1', 'QS-4 de Araguaína IV'), 2: ('RT ×2', 'LD Xambioá II - Araguanã - 2XB2'), 3: ('RT ×1', 'LD Araguaína - Babaçulândia - 2AR0'),
       4: ('RL ×4', 'LD UHE Isamu - Ponte Alta - 1UI1'), 5: ('RT ×3', 'LD UHE Isamu - Ponte Alta - 1UI1'), 6: ('RL ×1', 'QS-LD Palmas - Lajeado'),
       7: ('RL ×1', 'QS-8 de Palmas IV'), 8: ('RT ×1', 'LD Cristalândia - Lagoa da Confusão'), 9: ('RL ×1', 'QS-6 de Palmas IV'),
       10: ('RT ×1', 'LD Paraíso I - Caseara - 2PR0'), 11: ('RL ×3', 'QS-LD Palmas - Aparecida do Rio Negro'),
       12: ('RL ×1 + 3 TripSaver', 'QS-LD Gurupi - Duere'), 13: ('RT ×1', '1 de Figueirópolis - 2FGR')}
X0, X1 = 4007200, 11734495
y_h, rowh = 1330000, 285000
X = month_axis(macro, X0, X1, y_h, 300000, 4, 9, 3705000)
rect(macro, 457200, y_h, 3550000, 300000, NAVY)
text(macro, 457200, y_h, 3550000, 300000, 'Bloco · SIGCO · equipamento · projeto', 9, True, 'FFFFFF', 'l')
for i, n in enumerate(range(1, 14)):
    b = blocks[n]; r_ = b['rows']; y = 1630000 + i * rowh
    text(macro, 457200, y + 0, 3550000, rowh - 20000, f"B{n} · {b['sigco']} · {lab[n][0]} · {lab[n][1]}", 9, False, TXT, 'l')
    by = y + (rowh - 165000) / 2
    segs = [(r_['exec'][0], r_['exec'][1] + timedelta(days=1), ORANGE), (r_['fisc'][0], r_['cad'][1] + timedelta(days=1), TEAL),
            (r_['com'][0], r_['com'][1], DARKB if False else NAVY), (r_['et'][0], r_['ec'][1], GREY2)]
    for a, c, col in segs: rect(macro, X(a), by, max(X(c) - X(a), 15000), 165000, col)
    # meta de energização (fim do mês definido pela ASPO)
    mt = eom(b['meta'])
    rect(macro, X(mt) - 6000, y + 25000, 12000, rowh - 50000, RED)
    dmd = rect(macro, X(r_['op'][0]) - 70000, y + (rowh - 140000) / 2, 140000, 140000, GREEN, MSO_SHAPE.DIAMOND)
    lx = max(X(r_['ec'][1]), X(mt)) + 60000
    text(macro, lx, y, 1000000, rowh, f"Op. {dm(r_['op'][0])}", 8, True, RED if b['folga'] < 0 else NAVY, 'l')
ly = 5455000
leg = [(ORANGE, 'Execução'), (TEAL, 'Fiscalização e cadastro'), (NAVY, 'TCI / GAT / comissionamento'), (GREY2, 'Encerramentos')]
xx = 457200
for col, t in leg:
    rect(macro, xx, ly + 30000, 180000, 120000, col); text(macro, xx + 200000, ly, 1900000, 180000, t, 8, False, TXT, 'l'); xx += 2150000 if len(t) > 14 else 1100000
rect(macro, xx + 150000, ly + 20000, 140000, 140000, GREEN, MSO_SHAPE.DIAMOND); text(macro, xx + 330000, ly, 1500000, 180000, 'Equip. em operação', 8, False, TXT, 'l')
xx += 1950000
rect(macro, xx, ly + 10000, 12000, 160000, RED); text(macro, xx + 60000, ly, 2600000, 180000, 'Fim do mês de energização (ASPO)', 8, False, TXT, 'l')
text(macro, 457200, 5665000, 11277295, 520000,
     f'Equipamentos disponíveis somente após 05/04/2027: nenhuma execução inicia antes de 06/04/2027. Operação = início + 70 dias (execução de 30 d) ou + 85 dias (45 d). '
     f'Em vermelho: operação após o mês de energização definido pela ASPO (B{", B".join(str(n) for n in late[:-1])} e B{late[-1]}); encerramento contábil até {dmy(max(b["rows"]["ec"][1] for b in blocks.values()))} (limite 17/11/2027).',
     8, False, TXT, 'l', 't')


spec = {1: ('QS-4 de Araguaína IV', '1 RL 13,8 kV'), 2: ('LD Xambioá II - Araguanã - 2XB2', '2 RT 34,5 kV'), 3: ('LD Araguaína - Babaçulândia - 2AR0', '1 RT 34,5 kV'),
        4: ('LD UHE Isamu - Ponte Alta - 1UI1', '4 RL 34,5 kV'), 5: ('LD UHE Isamu - Ponte Alta - 1UI1', '3 RT 34,5 kV'), 6: ('QS-LD Palmas - Lajeado', '1 RL 34,5 kV'),
        7: ('QS-8 de Palmas IV', '1 RL 13,8 kV'), 8: ('LD Cristalândia - Lagoa da Confusão', '1 RT 34,5 kV'), 9: ('QS-6 de Palmas IV', '1 RL 13,8 kV'),
        10: ('LD Paraíso I - Caseara - 2PR0', '1 RT 34,5 kV'), 11: ('QS-LD Palmas - Aparecida do Rio Negro', '3 RL 34,5 kV'),
        12: ('QS-LD Gurupi - Duere', '1 RL rural 34,5 kV + 3 TripSaver'), 13: ('1 de Figueirópolis - 2FGR', '1 RT 13,8 kV + remanejamento de 1 RL')}

# ------------------------------------------------------------------ slides 11–13 — cronograma detalhado
def sig(b): return tuple((k, v[0], v[1]) for k, v in b['rows'].items())
groups = []
for n in range(1, 14):
    for g in groups:
        if sig(blocks[g[0]]) == sig(blocks[n]): g.append(n); break
    else: groups.append([n])
groups.sort(key=lambda g: (blocks[g[0]]['rows']['exec'][0], blocks[g[0]]['rows']['exec'][1]))
assert len(groups) == 6, groups
pages = [groups[0:2], groups[2:4], groups[4:6]]
# 6 grupos: (06/04 30d), (06/04 45d), (09/04), (16/04), (04/06), (11/06)
detail_slides = [S[10], S[11], S[12]]
for pg, sl in zip(pages, detail_slides):
    clear_body(sl)
    lo = min(blocks[g[0]]['rows']['exec'][0] for g in pg); hi = max(blocks[g[0]]['rows']['ec'][1] for g in pg)
    m_first, m_last = lo.month, hi.month
    X0d, X1d = 5757200, 11734495
    rect(sl, 457200, 1300000, 3650000, 300000, NAVY); text(sl, 457200, 1300000, 3650000, 300000, 'Tarefa', 9, True, 'FFFFFF', 'l')
    rect(sl, 4107200, 1300000, 1650000, 300000, NAVY); text(sl, 4107200, 1300000, 1650000, 300000, 'Dur. · Início – Término', 9, True, 'FFFFFF', 'c')
    Xd = month_axis(sl, X0d, X1d, 1300000, 300000, m_first, m_last, 0)
    pitch, bh = 235000, 135000
    y = 1600000
    for g in pg:
        b = blocks[g[0]]; r_ = b['rows']
        sig_txt = ', '.join(str(blocks[n]['sigco']) for n in g)
        if len(g) == 1:
            head = f"BLOCO {g[0]}  ·  SIGCO {sig_txt}  ·  {spec[g[0]][0]}  —  {spec[g[0]][1]}"
        else:
            head = 'BLOCOS ' + ', '.join(str(n) for n in g[:-1]) + f' e {g[-1]}  ·  SIGCO {sig_txt}  —  ' + ' | '.join(f'B{n}: {spec[n][1]}' for n in g)
        rect(sl, 457200, y, 11277295, pitch, 'E3E6F0'); text(sl, 457200, y, 11277295, pitch, head, 8, True, NAVY, 'l'); y += pitch
        tasks = [('Executar Obras', 'exec', ORANGE, 'd'), ('Fiscalização', 'fisc', TEAL, 'd'), ('Cadastro EO / Integração parcial', 'cad', TEAL2, 'd'),
                 ('Entrega do TCI / Integração final', 'tci', DARKB, 'd'), ('Abertura de chamado / Atendimento GAT', 'gat', LIME, 'd'),
                 ('Comissionar Equipamentos', 'com', NAVY, 'd'), ('Encerramento Técnico', 'et', GREY, 'du'), ('Encerramento Contábil', 'ec', GREY2, 'du')]
        for nm, k, col, un in tasks:
            a, c, dur = r_[k]
            text(sl, 457200, y, 3650000, pitch, '   ' + nm, 8, False, TXT, 'l')
            text(sl, 4107200, y, 1650000, pitch, f'{int(dur)} {un} · {dm(a)} – {dm(c)}', 8, False, TXT, 'c')
            xa = Xd(a); xc = Xd(c + timedelta(days=1)) if k not in ('et', 'ec') else Xd(c + timedelta(days=1))
            rect(sl, xa, y + (pitch - bh) / 2, max(xc - xa, 15000), bh, col)
            if k == 'exec': rect(sl, Xd(r_['conc'][0] + timedelta(days=1)) - 62500, y + (pitch - 125000) / 2, 125000, 125000, GREEN, MSO_SHAPE.DIAMOND)
            if k == 'com': rect(sl, Xd(r_['op'][0] + timedelta(days=1)) - 62500, y + (pitch - 125000) / 2, 125000, 125000, GREEN, MSO_SHAPE.DIAMOND)
            y += pitch
        rect(sl, 457200, y + 12000, 11277295, 6000, 'D9DCE8'); y += 40000
    ly = 6040000
    xx = 457200
    for col, t, w_ in [(ORANGE, 'Execução', 650000), (TEAL, 'Fiscalização', 800000), (TEAL2, 'Cadastro EO', 800000), (DARKB, 'TCI / Integração final', 1300000),
                       (LIME, 'GAT', 350000), (NAVY, 'Comissionamento', 1050000), (GREY, 'Enc. técnico', 800000), (GREY2, 'Enc. contábil', 800000)]:
        rect(sl, xx, ly + 30000, 180000, 120000, col); text(sl, xx + 200000, ly, w_, 180000, t, 8, False, TXT, 'l'); xx += 200000 + w_
    rect(sl, xx + 50000, ly + 15000, 150000, 150000, GREEN, MSO_SHAPE.DIAMOND); text(sl, xx + 230000, ly, 2800000, 180000, 'Marcos: obra concluída / em operação', 8, False, TXT, 'l')
for i, sl in enumerate(detail_slides, 1):
    set_runs(shp(sl, 'Title 2').text_frame.paragraphs[0], f'WBS – Cronograma Detalhado ({i}/3)')

# remove o 4º slide de detalhe
def remove_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId); lst.remove(sid); return
remove_slide(prs, S[13])

# ------------------------------------------------------------------ slide 10 — tabela WBS
tb = [s for s in S[9].shapes if s.has_table][0].table
for n in range(1, 14):
    row = tb.rows[n]
    vals = [f'1.{n}', f'BLOCO {n} – {spec[n][0]}', str(blocks[n]['sigco']), spec[n][1]]
    for c, v in enumerate(vals): cell_set(row.cells[c], v)

# ------------------------------------------------------------------ slide 6 — curva orçamentária
cs = S[5]
tbl = [s for s in cs.shapes if s.has_table][0].table
for i, cv in enumerate(curva, 1):
    row = tbl.rows[i]
    assert row.cells[0].text.strip() == str(cv['sigco']), (row.cells[0].text, cv['sigco'])
    for j, v in enumerate(cv['m']): cell_set(row.cells[3 + j], '-' if v < 0.5 else brl(v), '989898' if v < 0.5 else '4D4D4D')
    cell_set(row.cells[15], brl(cv['total']))
tr_ = tbl.rows[len(curva) + 1]
for j, v in enumerate(tot_m): cell_set(tr_.cells[3 + j], '-' if v < 0.5 else brl(v))
cell_set(tr_.cells[15], brl(tot))
ch = [s for s in cs.shapes if s.has_chart][0].chart
cd = CategoryChartData(); cd.categories = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
cd.add_series('Desembolso mensal (R$ mil)', [round(v / 1000, 1) for v in tot_m]); ch.replace_data(cd)
note = shp(cs, 'TextBox 10')
set_paras(note, [f'{acc_pct[3]*100:.0f}% do desembolso até abril/27, {acc_pct[4]*100:.0f}% até maio/27 e 100% até julho/27. Sem desembolso de jan a mar/27: execuções iniciam após 05/04/2027.',
                 'Distribuição em dias corridos entre o início da execução e o fim da fiscalização de cada bloco.'])
for pa in note.text_frame.paragraphs:
    for x in pa.runs: x.font.size = Pt(8.5)

# ------------------------------------------------------------------ slide 4 — resumo
rs = S[3]
tabs = {s.name: s.table for s in rs.shapes if s.has_table}
ti = tabs['Tabela interface']
com_first = min(b['rows']['com'][0] for b in blocks.values())
b3 = blocks[3]
ifc = [('Desligamento – 1ºs blocos', dmy(first_start), 'COI / DCMD'),
       ('Ajuste de proteção – 1º bloco', dmy(com_first), 'DEOP'),
       ('Comissionamento / GAT – 1º bloco', dmy(com_first), 'DMSL / GAT'),
       ('Energização – último bloco (B3)', dmy(b3['rows']['op'][0]), 'DCMD')]
for i, (a, d_, rsp) in enumerate(ifc, 1):
    cells = ti.rows[i].cells; cell_set(cells[1], a); cell_set(cells[2], d_); cell_set(cells[3], rsp)
tc = tabs['Tabela curva orçamentária']
for i, reg in enumerate(['Norte', 'Centro', 'Sul'], 1):
    vals = reg_m[reg]
    cells = tc.rows[i].cells
    for j in range(12):
        tx_ = '-' if vals[j] < 500 else brl(vals[j] / 1000)
        cell_set(cells[2 + j], tx_, size=5.5 if len(tx_) > 3 else None, tight=True)
    cell_set(cells[14], brl(vals[12] / 1000))
ta = tabs['Tabela aquisições']
cell_set(ta.rows[1].cells[2], dmy(first_start)); cell_set(ta.rows[2].cells[2], dmy(first_start)); cell_set(ta.rows[3].cells[2], dmy(first_start - timedelta(days=7)))
prem = shp(rs, 'Premissas lista')
ps = [p_.text for p_ in prem.text_frame.paragraphs]
ps[3] = 'Equipamentos disponíveis somente após 05/04/2027; nenhuma execução inicia antes de 06/04/2027.'
set_paras(prem, ps)
rst = shp(rs, 'Restrições lista')
rr_ = [p_.text for p_ in rst.text_frame.paragraphs]
rr_[1] = 'Energização no mês da ASPO; B1, B2, B4, B5 e B12 reprogramados para jun/27.'
set_paras(rst, rr_)
rk = shp(rs, 'Riscos lista')
set_paras(rk, ['Atraso no fornecimento de RL/RT e IEDs (lead time de até 210 dias – SEL).',
               'Energização fora do mês ASPO: B1, B2, B4, B5 e B12 (→ jun/27).',
               'Concentração: 11 blocos iniciam 06–16/04 e operam 15–30/06.',
               'Indisponibilidade de Linha Viva / cancelamento de desligamentos programados.',
               'Local sem validação prévia ou empreiteira sem capacidade executiva.',
               'Acidente de trabalho com paralisação da obra.'])

# ------------------------------------------------------------------ slide 5 — aquisições (necessário até = início da 1ª execução que usa o item)
aq = [s for s in S[4].shapes if s.has_table][0].table
start = {n: blocks[n]['rows']['exec'][0] for n in blocks}
def first(sigcos): return min(start[n] for n in blocks if blocks[n]['sigco'] in sigcos)
def firstn(ns): return min(start[n] for n in ns)
need = [first({91979, 92094, 92095, 90811}), first({92118, 92101, 92102}), first({89267, 70838, 70978, 90811, 89249}), first({89307}), first({91979}), first({89307}), first_start]
for i, d_ in enumerate(need, 1): cell_set(aq.rows[i].cells[3], dmy(d_))

# ------------------------------------------------------------------ slide 7 / 8 — restrição e risco de materiais
r7 = shp(S[6], 'Restrições lista')
r7p = [p_.text for p_ in r7.text_frame.paragraphs]
r7p[1] = 'Energização no mês da ASPO (abr, mai, jun e ago/2027); B1, B2, B4, B5 e B12 reprogramados para jun/27.'
set_paras(r7, r7p)
t3 = shp(S[7], 'Risco 3 – título')
set_runs(t3.text_frame.paragraphs[1], 'RL/RT/TripSaver e kit de periféricos disponíveis só após 05/04/2027 (lead time de até 210 dias)')

# ------------------------------------------------------------------ rodapés: numeração = posição real
slides = list(prs.slides)
for idx, sl in enumerate(slides, 1):
    for f in [s for s in sl.shapes if s.name == 'Rodapé']:
        set_runs(f.text_frame.paragraphs[0], f'{idx:02d}  ·  {NAME_UP}  ·  PN 2027')
prs.save(OUT)
print('saved', OUT, len(slides), 'slides; late blocks:', late)
for g in groups: print([ (n, blocks[n]['sigco']) for n in g], blocks[g[0]]['rows']['exec'][0], blocks[g[0]]['rows']['op'][0])
