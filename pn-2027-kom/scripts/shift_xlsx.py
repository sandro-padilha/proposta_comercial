import sys, copy, datetime as dt
import openpyxl
from openpyxl.comments import Comment
SRC, OUT = sys.argv[1], sys.argv[2]
LIMITE = dt.datetime(2027, 4, 5)           # equipamentos disponíveis somente após esta data
NOVO_INICIO = dt.datetime(2027, 4, 6)      # 1º dia após a disponibilidade
wb = openpyxl.load_workbook(SRC); wv = openpyxl.load_workbook(SRC, data_only=True)
w = wb['WBS_Cronograma']; v = wv['WBS_Cronograma']
op_rows = [r for r in range(1, w.max_row + 1) if isinstance(w.cell(r, 2).value, str) and w.cell(r, 2).value.strip() == 'Equipamentos em Operação']
changed = []
for r in op_rows:
    # início da execução = primeira linha de 'Executar Obras' acima da linha de operação
    ex = r - 7
    assert w.cell(ex, 2).value.strip() == 'Executar Obras', (r, w.cell(ex, 2).value)
    old_start, old_op = v.cell(ex, 6).value, v.cell(r, 6).value
    if old_start <= LIMITE:
        off = (old_op - old_start).days
        new_op = NOVO_INICIO + dt.timedelta(days=off)
        w.cell(r, 6).value = new_op
        w.cell(r, 6).comment = Comment(f'Reprogramado: início da execução passa de {old_start:%d/%m/%Y} para {NOVO_INICIO:%d/%m/%Y} (equipamentos disponíveis somente após 05/04/2027). Operação anterior: {old_op:%d/%m/%Y}.', 'Claude')
        changed.append((r, v.cell(ex - 0, 4).value if False else w.cell(ex, 4).value, old_start, old_op, new_op))
    # formato numérico da folga (estava em formato de data)
    w.cell(r, 10).number_format = '0'
# item 1.5: materiais disponíveis após 05/04/2027
r = 10
for c in range(1, 12):
    w.cell(r, c)._style = copy.copy(w.cell(9, c)._style)
w.cell(r, 5).value = 0
w.cell(r, 6).value = LIMITE
w.cell(r, 7).value = f'=F{r}+E{r}'
w.cell(r, 8).value = 'DEOP / Suprimentos'
w.cell(r, 11).value = 'A iniciar'
w.cell(r, 6).comment = Comment('Equipamentos (RL, RT e TripSaver) disponíveis somente após 05/04/2027: nenhuma execução inicia antes de 06/04/2027.', 'Claude')
rb = wb['Resumo_Blocos']
for rr in range(4, rb.max_row + 1): rb.cell(rr, 11).number_format = '0'
# premissa 5
pf = wb['Premissas_Fontes']
for rr in range(1, pf.max_row + 1):
    val = pf.cell(rr, 1).value
    if isinstance(val, str) and val.startswith('5. Equipamentos'):
        pf.cell(rr, 1).value = '5. Equipamentos (RL/RT/TripSaver) disponíveis somente após 05/04/2027: nenhuma execução de bloco inicia antes de 06/04/2027 (blocos antes programados para jan–abr/2027 foram reprogramados); O.S. da EPS emitida na semana anterior.'
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
red = Font(color='C00000', bold=True)
w.conditional_formatting.add(f'J5:J{w.max_row}', CellIsRule(operator='lessThan', formula=['0'], font=red))
rb.conditional_formatting.add(f'K4:K{rb.max_row}', CellIsRule(operator='lessThan', formula=['0'], font=red))
wb.save(OUT)
for c in changed: print(c)
