"""Gera WBS_Cronograma_PN2027.xlsx (EAP/WBS + cronograma + curva + aquisições) com fórmulas vivas."""
import sys
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from model import *

OUT = sys.argv[1]
P = build()

NAVY, ORANGE, TEAL, GREEN, GREY = '1F1F4C', 'F47B20', '009FC2', '628C0E', '989898'
F_BASE = 'Arial'
hdr_fill = PatternFill('solid', fgColor=NAVY)
sum_fill = PatternFill('solid', fgColor='E6E6E6')
blk_fill = PatternFill('solid', fgColor='DCE3F0')
thin = Side(style='thin', color='CCCCCC')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
DF = 'dd/mm/yyyy'

def font(**k):
    k.setdefault('name', F_BASE); k.setdefault('size', 10)
    return Font(**k)

def header(ws, row, labels, widths=None, height=30):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row, i, t)
        c.font = font(bold=True, color='FFFFFF'); c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = box
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w

wb = Workbook()

# ============================================================ Parâmetros
wp = wb.active; wp.title = 'Parâmetros'
wp['A1'] = 'Parâmetros do cronograma – PN 2027 (instalação de RL e RT)'; wp['A1'].font = font(bold=True, size=14, color=NAVY)
wp['A2'] = ('Prazos extraídos SOMENTE dos blocos de instalação de religadores (RL) e reguladores de tensão (RT) do WBS PN 2026. '
            'Células em azul são entradas editáveis; todo o cronograma recalcula a partir delas.')
wp['A2'].font = font(italic=True, color='4D4D4D'); wp['A2'].alignment = Alignment(wrap_text=True, vertical='top')
wp.merge_cells('A2:E2'); wp.row_dimensions[2].height = 42
header(wp, 4, ['Etapa (nome igual ao WBS 2026)', 'Dias', 'Tipo de dia', 'Regra de encadeamento', 'Origem do prazo (WBS PN 2026)'], [46, 8, 14, 52, 70])
rows = [
 ('Executar Obras – pacote até 4 equipamentos (P30)', 30, 'corrido', 'Início = dado de entrada (calculado de trás para frente a partir da meta de energização)', 'Blocos RT 72240, RT 71579, RT 72938, RL 81319 e AL SE Bielândia 71139 (1 a 4 equipamentos): 30 dias'),
 ('Executar Obras – pacote com 5 ou mais equipamentos (P45)', 45, 'corrido', 'Idem', 'BLOCO 1 (5 RL + 1 remanej., SIGCO 71018) e BLOCO 3 (5 RL, SIGCO 73521): 45 dias'),
 ('Fiscalização', 12, 'corrido', 'Inicia no 1º dia útil após Obras Concluídas', 'Todos os blocos RL/RT: 12 dias'),
 ('Cadastro EO/ Integração parcial', 10, 'corrido', 'Inicia no 1º dia útil após a Fiscalização', 'Todos os blocos RL/RT: 10 dias'),
 ('Entrega do TCI pela empreiteira / Integração final', 15, 'corrido', 'Inicia no 1º dia útil após o Cadastro EO', 'Todos os blocos RL/RT: 15 dias'),
 ('Abertura de chamado/ Atendimento GAT', 8, 'corrido', 'Paralela: inicia junto com a Entrega do TCI', 'Todos os blocos RL/RT: 8 dias'),
 ('Comissionar Equipamentos', 15, 'corrido', 'Paralela: inicia junto com a Entrega do TCI; término = Equipamentos em Operação', 'Todos os blocos RL/RT: 15 dias'),
 ('Encerramento Técnico', 7, 'útil', 'Inicia no 1º dia útil após Equipamentos em Operação', 'Todos os blocos RL/RT: 7 dias úteis'),
 ('Encerramento Contábil', 3, 'útil', 'Inicia no 1º dia útil após o Encerramento Técnico', 'Todos os blocos RL/RT: 3 dias úteis'),
]
for i, r in enumerate(rows, 5):
    for j, v in enumerate(r, 1):
        c = wp.cell(i, j, v); c.border = box; c.font = font()
        c.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center' if j in (2, 3) else 'left')
    wp.cell(i, 2).font = font(color='0000FF', bold=True)
    wp.row_dimensions[i].height = 30
PR = dict(p30='Parâmetros!$B$5', p45='Parâmetros!$B$6', fisc='Parâmetros!$B$7', cad='Parâmetros!$B$8', tci='Parâmetros!$B$9',
          gat='Parâmetros!$B$10', com='Parâmetros!$B$11', et='Parâmetros!$B$12', ec='Parâmetros!$B$13')
wp['A15'] = 'Parcela de material na curva orçamentária'; wp['A15'].font = font(bold=True)
wp['B15'] = MAT_PCT; wp['B15'].number_format = '0%'; wp['B15'].font = font(color='0000FF', bold=True)
wp['D15'] = 'Referência: KOM 2026 (material R$ 15,48 mi de R$ 25,79 mi = 60%). Material reconhecido no mês de início da execução; serviço rateado pelos dias de execução.'
wp['D15'].font = font(italic=True, color='4D4D4D'); wp['D15'].alignment = Alignment(wrap_text=True, vertical='top'); wp.merge_cells('D15:E15'); wp.row_dimensions[15].height = 30
wp['A16'] = 'Restrição: encerramento contábil até'; wp['A16'].font = font(bold=True)
wp['B16'] = date(2027, 11, 17); wp['B16'].number_format = DF; wp['B16'].font = font(color='0000FF', bold=True)
wp['D16'] = 'Template KOM PMO-ETO (restrição padrão).'; wp['D16'].font = font(italic=True, color='4D4D4D')
wp['A18'] = 'Feriados considerados (dias úteis)'; wp['A18'].font = font(bold=True)
hol_names = ['Independência', 'N. Sra. Aparecida', 'Finados', 'Proclamação da República', 'Natal', 'Confraternização Universal',
             'Sexta-feira Santa', 'Tiradentes', 'Dia do Trabalho', 'Corpus Christi', 'Independência', 'N. Sra. Aparecida', 'Finados',
             'Proclamação da República', 'Natal']
for i, (h, n) in enumerate(zip(HOLIDAYS, hol_names), 19):
    wp.cell(i, 1, n).font = font(); c = wp.cell(i, 2, h); c.number_format = DF; c.font = font(color='0000FF')
HOL_RNG = f"Parâmetros!$B$19:$B${18 + len(HOLIDAYS)}"
wb.defined_names['Feriados'] = DefinedName('Feriados', attr_text=HOL_RNG)
wp.freeze_panes = 'A5'

# ============================================================ Carteira
wc = wb.create_sheet('Carteira')
wc['A1'] = 'Carteira PN 2027 – 12 SIGCOs de instalação de RL e RT (GP: Sandro Padilha Dias Junior – DCMD)'; wc['A1'].font = font(bold=True, size=14, color=NAVY)
header(wc, 3, ['SIGCO', 'Capítulo PN', 'Projeto', 'Escopo (equipamentos)', 'RL', 'RT', 'TripSaver', 'Nº equip.', 'Pacote de prazo', 'Meta de energização',
               'Orçamento Agosto (R$)', 'Orçamento 08.10 (R$) – aprovado', 'Δ 08.10 − Agosto (R$)', 'Benefício esperado'],
       [9, 18, 52, 46, 5, 5, 9, 9, 12, 14, 18, 20, 18, 38], 42)
CR = {}
for i, p in enumerate(P, 4):
    ts = 3 if p['sigco'] == 91979 else 0
    vals = [p['sigco'], p['cap'], p['nome'], p['eq'], p['rl'], p['rt'], ts, p['n'], p['pkg'], eom(*p['meta']), p['b_ago'], p['b_1010'], f'=L{i}-K{i}', p['ben']]
    for j, v in enumerate(vals, 1):
        c = wc.cell(i, j, v); c.font = font(); c.border = box; c.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center' if j in (1, 5, 6, 7, 8, 9, 10) else 'left')
    wc.cell(i, 10).number_format = DF
    for j in (11, 12, 13): wc.cell(i, j).number_format = '#,##0.00'
    wc.cell(i, 10).font = font(color='0000FF'); wc.cell(i, 9).font = font(color='0000FF'); wc.cell(i, 12).font = font(color='0000FF', bold=True); wc.cell(i, 11).font = font(color='0000FF')
    wc.row_dimensions[i].height = 30
    CR[p['sigco']] = i
tr = 4 + len(P)
wc.cell(tr, 3, 'Total da carteira').font = font(bold=True)
for j, col in ((5, 'E'), (6, 'F'), (7, 'G'), (8, 'H'), (11, 'K'), (12, 'L'), (13, 'M')):
    c = wc.cell(tr, j, f'=SUM({col}4:{col}{tr-1})'); c.font = font(bold=True); c.fill = sum_fill; c.border = box
    if j >= 11: c.number_format = '#,##0.00'
wc.cell(tr + 2, 1, 'Fontes: Planejamento PN.xlsx (aba CAPEX), Cadastro_CAPEX_2027_2029 (Orçamento 08.10), fichas dos projetos 2027 (PDF), TM-Transfer Meeting 06/10/2026 (meta de energização).').font = font(italic=True, color='4D4D4D')
wc.cell(tr + 3, 1, 'RL = religador (inclui 1 RL rural 34,5 kV no SIGCO 91979 e 1 RL remanejado no 89307); RT = regulador de tensão; pacote P30 = até 4 equipamentos, P45 = 5 ou mais.').font = font(italic=True, color='4D4D4D')
wc.freeze_panes = 'D4'

# ============================================================ WBS (formato Plan1 + EAP)
ww = wb.create_sheet('WBS', 0)
header(ww, 1, ['Nome da tarefa', 'Duração', 'Início', 'Término', 'Código EAP', 'Tipo', 'Predecessora', 'Responsável (função)', 'SIGCO', 'Base do prazo'],
       [78, 9, 12, 12, 11, 14, 24, 30, 9, 58], 28)
ww.sheet_properties.outlinePr.summaryBelow = False
R = {}          # R[sigco][key] = linha
row = 2

def put(r, name, B=None, C=None, D=None, code='', tipo='', pred='', resp='', sig=None, base='', level=0, fill=None, bold=False, key=None):
    ww.cell(r, 1, name); ww.cell(r, 2, B); ww.cell(r, 3, C); ww.cell(r, 4, D); ww.cell(r, 5, code); ww.cell(r, 6, tipo)
    ww.cell(r, 7, pred); ww.cell(r, 8, resp); ww.cell(r, 9, sig); ww.cell(r, 10, base)
    for j in range(1, 11):
        c = ww.cell(r, j); c.font = font(bold=bold); c.border = box
        c.alignment = Alignment(vertical='center', horizontal='left' if j in (1, 7, 8, 10) else 'center', indent=level if j == 1 else 0, wrap_text=(j in (7, 8, 10)))
        if fill: c.fill = fill
    ww.cell(r, 3).number_format = DF; ww.cell(r, 4).number_format = DF
    ww.cell(r, 2).number_format = '0'

# raiz
root = row; put(root, 'PN 2027 – Instalação de Equipamentos Especiais (RL e RT) – 12 SIGCOs', code='1', tipo='Projeto', resp='GP Sandro Padilha (DCMD)', fill=sum_fill, bold=True); row += 1
# 1.1 Gerenciamento
g = row; put(g, 'Gerenciamento do Projeto (PMBOK)', code='1.1', tipo='Entrega', resp='GP / ASGP (PMO)', level=1, fill=blk_fill, bold=True); row += 1
pmo = [('Transfer Meeting – TM CAPEX (realizado)', date(2026, 10, 6), '1.1.1', 'Iniciação', 'ASGP / ASPO', 'Calendário PMO (ATA TM 06/10/2026)'),
       ('Kickoff Meeting – KOM CAPEX', date(2026, 10, 29), '1.1.2', 'Planejamento', 'GP / PMO', 'Calendário PMO (ATA TM 06/10/2026)'),
       ('Project Review Meeting – PRM CAPEX (consolidar planejamento / salvar linha de base)', date(2026, 12, 2), '1.1.3', 'Planejamento', 'GP / PMO', 'Calendário PMO (ATA TM 06/10/2026)')]
gm = []
for nm, d, cd, tp, rs, bs in pmo:
    put(row, nm, 0, d, f'=C{row}+B{row}', cd, 'Marco', '', rs, None, bs, level=2); gm.append(row); row += 1
mon = row  # monitoramento e controle (derivado)
put(mon, 'Monitoramento e controle (FUM mensal, reports, forecast e gestão de riscos)', None, None, None, '1.1.4', 'Pacote', f'PRM', 'GP', None, 'Do PRM ao último Encerramento Contábil (duração derivada)', level=2); row += 1
close_row = row
put(close_row, 'Encerramento do projeto: Close Meeting (CM), termo de encerramento, lições aprendidas', None, None, None, '1.1.5', 'Pacote', 'Último Enc. Contábil', 'GP / PMO', None, 'Após o último Encerramento Contábil (restrição: até 17/11/2027)', level=2); row += 1
g_end = row - 1

# 1.2 Aquisições
a = row; put(a, 'Aquisições e Contratações', code='1.2', tipo='Entrega', resp='DCMD / Suprimentos', level=1, fill=blk_fill, bold=True); row += 1
acq_rows = {}
for cd, nm, rs in (('1.2.1', 'Equipamentos RL/RT disponíveis para a 1ª frente (PMA/RFP – data necessária em campo)', 'DCMD (PMA de materiais)'),
                   ('1.2.2', 'Contratação da empreiteira executora (EPS) – O.S. assinada para a 1ª frente', 'DCMD (contratação de obras)')):
    put(row, nm, 0, None, f'=C{row}+B{row}', cd, 'Marco', '', rs, None, 'Data necessária = início da 1ª execução (sem prazo de aquisição adotado: validar lead time com Suprimentos)', level=2)
    acq_rows[cd] = row; row += 1

# 1.3.x blocos por SIGCO
blk_hdr = row; put(blk_hdr, 'Instalação, comissionamento e encerramento por SIGCO', code='1.3', tipo='Entrega', resp='DCMD (COCM) / DEOP / DMSL', level=1, fill=blk_fill, bold=True); row += 1
names = dict(PHASES)
resp = dict(exec='EPS (executora) / Supervisor de Obras DCMD', conc='EPS', fisc='Fiscalização de obras (Resulta)', cad='DCMD (Cadastro EO/TS)',
            tci='EPS / DCMD', gat='DCMD / GAT', com='DMSL / DEOP (ajuste de proteção)', op='DMSL / DEOP', et='DCMD (Encerramento Técnico)', ec='DCMD / Contabilidade')
basis = dict(exec='PN 2026 – blocos RL/RT', conc='Marco', fisc='PN 2026 – blocos RL/RT', cad='PN 2026 – blocos RL/RT', tci='PN 2026 – blocos RL/RT',
             gat='PN 2026 – blocos RL/RT', com='PN 2026 – blocos RL/RT', op='Marco = término do comissionamento', et='PN 2026 – blocos RL/RT (dias úteis)', ec='PN 2026 – blocos RL/RT (dias úteis)')
pred = dict(exec='Início por meta de energização', conc='Executar Obras', fisc='Obras Concluídas', cad='Fiscalização', tci='Cadastro EO',
            gat='TCI (início)', com='TCI (início)', op='Comissionar Equip.', et='Equip. em Operação', ec='Encerramento Técnico')
for k, p in enumerate(P, 1):
    s = p['sigco']; R[s] = {}
    b = row; R[s]['blk'] = b
    put(b, f"{p['nome']} ({s})", None, None, None, f'1.3.{k}', 'Pacote', '', 'GP Sandro Padilha', s, p['eq'], level=2, fill=sum_fill, bold=True); row += 1
    for n, (key, nm) in enumerate(PHASES, 1):
        R[s][key] = row
        tipo = 'Marco' if key in ('conc', 'op') else 'Atividade'
        put(row, nm, None, None, None, f'1.3.{k}.{n}', tipo, pred[key], resp[key], s, basis[key], level=3); row += 1
    r = R[s]
    pk = PR['p30'] if p['pkg'] == 'P30' else PR['p45']
    ww.cell(r['exec'], 2, f'={pk}'); ww.cell(r['exec'], 3, p['S']); ww.cell(r['exec'], 3).font = font(color='0000FF', bold=True); ww.cell(r['exec'], 4, f"=C{r['exec']}+B{r['exec']}")
    ww.cell(r['conc'], 2, 0); ww.cell(r['conc'], 3, f"=D{r['exec']}"); ww.cell(r['conc'], 4, f"=C{r['conc']}+B{r['conc']}")
    ww.cell(r['fisc'], 2, f"={PR['fisc']}"); ww.cell(r['fisc'], 3, f"=WORKDAY(D{r['conc']},1,Feriados)"); ww.cell(r['fisc'], 4, f"=C{r['fisc']}+B{r['fisc']}")
    ww.cell(r['cad'], 2, f"={PR['cad']}"); ww.cell(r['cad'], 3, f"=WORKDAY(D{r['fisc']},1,Feriados)"); ww.cell(r['cad'], 4, f"=C{r['cad']}+B{r['cad']}")
    ww.cell(r['tci'], 2, f"={PR['tci']}"); ww.cell(r['tci'], 3, f"=WORKDAY(D{r['cad']},1,Feriados)"); ww.cell(r['tci'], 4, f"=C{r['tci']}+B{r['tci']}")
    ww.cell(r['gat'], 2, f"={PR['gat']}"); ww.cell(r['gat'], 3, f"=C{r['tci']}"); ww.cell(r['gat'], 4, f"=C{r['gat']}+B{r['gat']}")
    ww.cell(r['com'], 2, f"={PR['com']}"); ww.cell(r['com'], 3, f"=C{r['tci']}"); ww.cell(r['com'], 4, f"=C{r['com']}+B{r['com']}")
    ww.cell(r['op'], 2, 0); ww.cell(r['op'], 3, f"=D{r['com']}"); ww.cell(r['op'], 4, f"=C{r['op']}+B{r['op']}")
    ww.cell(r['et'], 2, f"={PR['et']}"); ww.cell(r['et'], 3, f"=WORKDAY(D{r['op']},1,Feriados)"); ww.cell(r['et'], 4, f"=WORKDAY(C{r['et']},B{r['et']},Feriados)")
    ww.cell(r['ec'], 2, f"={PR['ec']}"); ww.cell(r['ec'], 3, f"=WORKDAY(D{r['et']},1,Feriados)"); ww.cell(r['ec'], 4, f"=WORKDAY(C{r['ec']},B{r['ec']},Feriados)")
    first, last = r['exec'], r['ec']
    ww.cell(b, 3, f'=MIN(C{first}:C{last})'); ww.cell(b, 4, f'=MAX(D{first}:D{last})'); ww.cell(b, 2, f'=D{b}-C{b}')
    for rr in range(first, last + 1): ww.row_dimensions[rr].outlineLevel = 1
blk_last = row - 1
first_exec = [R[p['sigco']]['exec'] for p in P]
last_ec = [R[p['sigco']]['ec'] for p in P]
# marcos de aquisição e monitoramento (fórmulas)
exec_starts = ','.join(f'C{r}' for r in first_exec)
ec_ends = ','.join(f'D{r}' for r in last_ec)
for cd, rr in acq_rows.items():
    ww.cell(rr, 3, f'=MIN({exec_starts})'); ww.cell(rr, 4, f'=C{rr}+B{rr}')
ww.cell(mon, 3, f'=D{gm[2]}'); ww.cell(mon, 4, f'=MAX({ec_ends})'); ww.cell(mon, 2, f'=D{mon}-C{mon}')
ww.cell(close_row, 3, f'=MAX({ec_ends})'); ww.cell(close_row, 4, f'=Parâmetros!$B$16'); ww.cell(close_row, 2, f'=D{close_row}-C{close_row}')
ww.cell(g, 3, f'=MIN(C{gm[0]}:C{close_row})'); ww.cell(g, 4, f'=MAX(D{gm[0]}:D{close_row})'); ww.cell(g, 2, f'=D{g}-C{g}')
ww.cell(a, 3, f'=MIN(C{a+1}:C{a+2})'); ww.cell(a, 4, f'=MAX(D{a+1}:D{a+2})'); ww.cell(a, 2, f'=D{a}-C{a}')
ww.cell(blk_hdr, 3, f'=MIN(C{blk_hdr+1}:C{blk_last})'); ww.cell(blk_hdr, 4, f'=MAX(D{blk_hdr+1}:D{blk_last})'); ww.cell(blk_hdr, 2, f'=D{blk_hdr}-C{blk_hdr}')
ww.cell(root, 3, f'=MIN(C{g}:C{blk_last})'); ww.cell(root, 4, f'=MAX(D{g}:D{blk_last})'); ww.cell(root, 2, f'=D{root}-C{root}')
for rr in (g, a, blk_hdr): pass
for rr in range(1, blk_last + 1):
    for j in (3, 4): ww.cell(rr, j).number_format = DF
ww.freeze_panes = 'B2'
ww.auto_filter.ref = f'A1:J{blk_last}'
ww.cell(blk_last + 2, 1, 'Datas em azul = entradas (início da execução, calculado de trás para frente a partir da meta de energização). Demais datas = fórmulas (WORKDAY com feriados da aba Parâmetros).').font = font(italic=True, color='4D4D4D')
ww.cell(blk_last + 3, 1, 'Escopo da EAP: 100% do trabalho do PN 2027 de RL/RT; só as atividades de instalação têm prazo (histórico PN 2026 de RL/RT). Fornecimento/contratação aparecem como marcos de "data necessária".').font = font(italic=True, color='4D4D4D')

# ============================================================ Macro
wm = wb.create_sheet('Cronograma Macro', 1)
wm['A1'] = 'Cronograma macro – PN 2027: instalação de RL e RT'; wm['A1'].font = font(bold=True, size=14, color=NAVY)
header(wm, 3, ['SIGCO', 'Projeto', 'Pacote', 'Equip.', 'Início da execução', 'Obras concluídas', 'Início comissionamento', 'Equip. em operação', 'Meta de energização',
               'Folga vs. meta (dias)', 'Fim Enc. Técnico', 'Fim Enc. Contábil', 'Duração total (dias)', 'Folga p/ limite contábil 17/11/2027 (dias)'],
       [9, 44, 8, 8, 14, 14, 14, 14, 14, 12, 14, 14, 12, 18], 44)
for i, p in enumerate(P, 4):
    s = p['sigco']; r = R[s]
    vals = [s, f"=Carteira!C{CR[s]}", f"=Carteira!I{CR[s]}", f"=Carteira!H{CR[s]}", f"=WBS!C{r['exec']}", f"=WBS!D{r['conc']}", f"=WBS!C{r['com']}", f"=WBS!D{r['op']}",
            f"=Carteira!J{CR[s]}", f"=I{i}-H{i}", f"=WBS!D{r['et']}", f"=WBS!D{r['ec']}", f"=L{i}-E{i}", f"=Parâmetros!$B$16-L{i}"]
    for j, v in enumerate(vals, 1):
        c = wm.cell(i, j, v); c.font = font(); c.border = box
        c.alignment = Alignment(vertical='center', wrap_text=True, horizontal='left' if j == 2 else 'center')
        if j in (5, 6, 7, 8, 9, 11, 12): c.number_format = DF
    wm.row_dimensions[i].height = 28
n = 3 + len(P)
wm.cell(n + 2, 1, 'Equipamentos em operação ≤ meta de energização do mês (TM 06/10/2026). Folga = 0 indica cronograma sem reserva (agosto/2027).').font = font(italic=True, color='4D4D4D')
wm.freeze_panes = 'C4'

# ============================================================ Gantt
wg = wb.create_sheet('Gantt', 2)
wg['A1'] = 'Gantt semanal – PN 2027 (barras recalculam com as datas da aba WBS)'; wg['A1'].font = font(bold=True, size=14, color=NAVY)
wg['E2'] = 'Legenda:'; wg['E2'].font = font(bold=True, size=8)
leg = [('Execução', NAVY), ('Fiscalização / Cadastro / TCI', TEAL), ('GAT / Comissionamento', GREEN), ('Encerramento', ORANGE), ('Marco', 'C00000')]
for i, (t, col) in enumerate(leg):
    c0 = 6 + i * 9
    wg.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 8)
    c = wg.cell(2, c0, t)
    c.font = font(bold=True, color='FFFFFF', size=8); c.fill = PatternFill('solid', fgColor=col); c.alignment = Alignment(horizontal='center', vertical='center')
HR = 4  # linha de cabeçalho (datas das semanas)
w0 = date(2026, 9, 28)
nweeks = 61
FC = 6
for j, t in enumerate(['Código EAP', 'Tarefa', 'Início', 'Término', 'tipo'], 1):
    c = wg.cell(HR, j, t); c.font = font(bold=True, color='FFFFFF'); c.fill = hdr_fill; c.alignment = Alignment(horizontal='center', vertical='center'); c.border = box
    wg.cell(HR - 1, j).fill = hdr_fill
for j, w in enumerate([11, 70, 11, 11, 5], 1): wg.column_dimensions[L(j)].width = w
for w in range(nweeks):
    d = w0 + timedelta(days=7 * w)
    col = FC + w
    c = wg.cell(HR, col, d); c.number_format = 'dd/mm'; c.font = font(size=7, bold=True, color='FFFFFF'); c.fill = hdr_fill
    c.alignment = Alignment(text_rotation=90, horizontal='center', vertical='center')
    wg.column_dimensions[L(col)].width = 2.6
    m = wg.cell(HR - 1, col)
    if w == 0 or (w0 + timedelta(days=7 * (w - 1))).month != d.month:
        m.value = d.strftime('%m/%y'); m.font = font(size=8, bold=True, color='FFFFFF')
    m.fill = hdr_fill
wg.row_dimensions[HR].height = 38
keymap = {}
for s in R:
    for k in R[s]:
        keymap[R[s][k]] = k
gr = HR + 1
for rr in range(2, blk_last + 1):
    nm = ww.cell(rr, 1).value
    if nm is None: continue
    typ = ww.cell(rr, 6).value
    wg.cell(gr, 1, f'=WBS!E{rr}'); wg.cell(gr, 2, f'=WBS!A{rr}'); wg.cell(gr, 3, f'=WBS!C{rr}'); wg.cell(gr, 4, f'=WBS!D{rr}')
    lvl = ww.cell(rr, 1).alignment.indent
    key = keymap.get(rr)
    if key == 'exec': kk = 'exec'
    elif key in ('fisc', 'cad', 'tci'): kk = 'adm'
    elif key in ('gat', 'com'): kk = 'com'
    elif key in ('et', 'ec'): kk = 'enc'
    elif key in ('conc', 'op') or typ == 'Marco': kk = 'ms'
    else: kk = 'sum'
    wg.cell(gr, 5, kk).font = font(size=7, color='BBBBBB')
    for j in range(1, 5):
        c = wg.cell(gr, j); c.font = font(size=9, bold=(kk == 'sum' and lvl <= 2)); c.border = box
        c.alignment = Alignment(vertical='center', indent=lvl if j == 2 else 0, horizontal='left' if j == 2 else 'center')
        if j in (3, 4): c.number_format = DF
        if kk == 'sum' and lvl <= 2: c.fill = sum_fill
    gr += 1
g_last = gr - 1
rng = f'{L(FC)}{HR+1}:{L(FC+nweeks-1)}{g_last}'
cl = L(FC)
for kk, col in (('exec', NAVY), ('adm', TEAL), ('com', GREEN), ('enc', ORANGE), ('ms', 'C00000')):
    wg.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($E{HR+1}="{kk}",{cl}${HR}<=$D{HR+1},{cl}${HR}+6>=$C{HR+1})'], fill=PatternFill('solid', bgColor=col, fgColor=col), stopIfTrue=True))
wg.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($E{HR+1}="sum",{cl}${HR}<=$D{HR+1},{cl}${HR}+6>=$C{HR+1})'], fill=PatternFill('solid', bgColor='9A9AB8', fgColor='9A9AB8'), stopIfTrue=True))
wg.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(Parâmetros!$B$16>={cl}${HR},Parâmetros!$B$16<={cl}${HR}+6)'], fill=PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')))
wg.cell(g_last + 2, 2, 'Coluna rosada = semana do limite de encerramento contábil (17/11/2027).').font = font(italic=True, color='4D4D4D', size=9)
wg.freeze_panes = wg.cell(HR + 1, FC)

# ============================================================ Curva orçamentária
wq = wb.create_sheet('Curva orçamentária', 3)
wq['A1'] = 'Curva orçamentária 2027 alinhada à WBS (R$) – proposta para atualização do SIACO'; wq['A1'].font = font(bold=True, size=14, color=NAVY)
wq['A2'] = ('Material no mês de início da execução (parâmetro 60%) e serviço rateado pelos dias de execução. Total = Orçamento 08.10 (aprovado). '
            'O cadastro 08.10 distribui o valor linearmente de jan a dez; como no KOM 2026, recomenda-se ajustar o SIACO para coincidir com a curva da WBS.')
wq['A2'].font = font(italic=True, color='4D4D4D'); wq['A2'].alignment = Alignment(wrap_text=True, vertical='top'); wq.merge_cells('A2:P2'); wq.row_dimensions[2].height = 40
months = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
header(wq, 4, ['SIGCO', 'Projeto', 'Total aprovado (08.10)'] + months + ['Total curva', 'Check (Σ = total)'], [9, 42, 16] + [11] * 12 + [14, 12], 30)
for i, p in enumerate(P, 5):
    s = p['sigco']; r = R[s]
    wq.cell(i, 1, s); wq.cell(i, 2, f'=Carteira!C{CR[s]}'); wq.cell(i, 3, f'=Carteira!L{CR[s]}')
    S_, E_, D_ = f"WBS!$C${r['exec']}", f"WBS!$D${r['exec']}", f"WBS!$B${r['exec']}"
    for m in range(1, 13):
        col = 3 + m
        ms = f'DATE(2027,{m},1)'
        f = (f"=$C{i}*(Parâmetros!$B$15*(YEAR({S_})=2027)*(MONTH({S_})={m})"
             f"+(1-Parâmetros!$B$15)*MAX(0,MIN({E_},EDATE({ms},1))-MAX({S_},{ms}))/{D_})")
        wq.cell(i, col, f)
    wq.cell(i, 16, f'=SUM(D{i}:O{i})'); wq.cell(i, 17, f'=IF(ABS(P{i}-C{i})<0.5,"ok","DIVERGE")')
    for j in range(1, 18):
        c = wq.cell(i, j); c.font = font(); c.border = box
        c.alignment = Alignment(vertical='center', horizontal='left' if j == 2 else ('center' if j in (1, 17) else 'right'))
        if 3 <= j <= 16: c.number_format = '#,##0'
nq = 4 + len(P)
wq.cell(nq + 1, 2, 'Total mensal – curva WBS').font = font(bold=True)
for j in range(3, 17):
    c = wq.cell(nq + 1, j, f'=SUM({L(j)}5:{L(j)}{nq})'); c.font = font(bold=True); c.fill = sum_fill; c.number_format = '#,##0'; c.border = box
wq.cell(nq + 2, 2, 'Acumulado – curva WBS').font = font(bold=True)
for m in range(1, 13):
    j = 3 + m
    c = wq.cell(nq + 2, j, f'=SUM($D{nq+1}:{L(j)}{nq+1})'); c.number_format = '#,##0'; c.font = font()
# cadastro linear
import openpyxl
cad = openpyxl.load_workbook('/tmp/claude-0/-home-user-proposta-comercial/9da2235d-c067-5426-823d-b6e1ef8d168f/scratchpad/zip/Cadastro_CAPEX_2027_2029_24 08 2026_PMO.xlsx', data_only=True)['Mensalização Capex 08.10']
lin = {}
for r_ in cad.iter_rows(min_row=2, values_only=True):
    if r_[1] in {p['sigco'] for p in P} and r_[3] == 'DCMD': lin[r_[1]] = list(r_[4:16])
base = nq + 5
wq.cell(base - 1, 1, 'Referência: cadastro CAPEX "Mensalização 08.10" (rateio linear)').font = font(bold=True, color=NAVY)
header(wq, base, ['SIGCO', 'Projeto', 'Total'] + months + ['Total', ''], None, 24)
for i, p in enumerate(P, base + 1):
    s = p['sigco']; wq.cell(i, 1, s); wq.cell(i, 2, f'=Carteira!C{CR[s]}'); wq.cell(i, 3, f'=SUM(D{i}:O{i})')
    for m, v in enumerate(lin[s], 4): wq.cell(i, m, v)
    wq.cell(i, 16, f'=SUM(D{i}:O{i})')
    for j in range(1, 17):
        c = wq.cell(i, j); c.font = font(color='4D4D4D'); c.border = box
        c.alignment = Alignment(horizontal='left' if j == 2 else ('center' if j == 1 else 'right'))
        if j >= 3: c.number_format = '#,##0'
nl = base + len(P)
wq.cell(nl + 1, 2, 'Total mensal – cadastro 08.10').font = font(bold=True)
for j in range(3, 17):
    c = wq.cell(nl + 1, j, f'=SUM({L(j)}{base+1}:{L(j)}{nl})'); c.font = font(bold=True); c.fill = sum_fill; c.number_format = '#,##0'; c.border = box
wq.cell(nl + 2, 2, 'Diferença (curva WBS − cadastro)').font = font(bold=True)
for j in range(4, 16):
    c = wq.cell(nl + 2, j, f'={L(j)}{nq+1}-{L(j)}{nl+1}'); c.number_format = '#,##0;[Red]-#,##0'; c.font = font()
wq.freeze_panes = 'C5'

# ============================================================ Aquisições
wa = wb.create_sheet('Aquisições', 4)
wa['A1'] = 'Equipamentos e serviços a adquirir/contratar – data necessária em campo'; wa['A1'].font = font(bold=True, size=14, color=NAVY)
header(wa, 3, ['SIGCO', 'Item', 'Tensão', 'Quantidade', 'Data necessária (início da execução)', 'Nº PMA/RFP', 'Status'], [9, 56, 10, 11, 22, 14, 18], 36)
items = [
 (90811, 'Regulador de tensão', '34,5 kV', 3), (90811, 'Religador automático', '34,5 kV', 4), (90811, 'Chave faca (1 N.F + 1 N.A) + poste e 0,25 km de RD MT cabo 3#1/0 AWG', '34,5 kV', 1),
 (91979, 'Religador rural', '34,5 kV', 1), (91979, 'TripSaver (religador monofásico)', '—', 3),
 (92118, 'Religador automático', '13,8 kV', 1),
 (89267, 'Regulador de tensão', '34,5 kV', 2),
 (92094, 'Religador automático', '34,5 kV', 3),
 (92095, 'Religador automático', '34,5 kV', 1),
 (92101, 'Religador automático (N.A)', '13,8 kV', 1),
 (92102, 'Religador automático (N.F)', '13,8 kV', 1),
 (70838, 'Regulador de tensão', '34,5 kV', 1),
 (70978, 'Regulador de tensão', '34,5 kV', 1),
 (89249, 'Regulador de tensão', '34,5 kV', 1),
 (89307, 'Regulador de tensão', '13,8 kV', 1), (89307, 'Religador 13,8 kV – remanejamento (equipamento existente)', '13,8 kV', 1),
]
for i, (s, it, tv, q) in enumerate(items, 4):
    wa.cell(i, 1, s); wa.cell(i, 2, it); wa.cell(i, 3, tv); wa.cell(i, 4, q); wa.cell(i, 5, f"=WBS!C{R[s]['exec']}"); wa.cell(i, 6, ''); wa.cell(i, 7, 'A confirmar')
    for j in range(1, 8):
        c = wa.cell(i, j); c.font = font(); c.border = box; c.alignment = Alignment(vertical='center', horizontal='left' if j == 2 else 'center', wrap_text=True)
    wa.cell(i, 5).number_format = DF
ni = 3 + len(items)
wa.cell(ni + 2, 1, 'Resumo: 9 RT (8 de 34,5 kV e 1 de 13,8 kV); 12 RL novos (9 de 34,5 kV e 3 de 13,8 kV) + 1 RL remanejado; 3 TripSavers.').font = font(bold=True)
wa.cell(ni + 3, 1, 'Atenção: o prazo de fornecimento não foi usado no cronograma. A ATA do TM (06/10/2026) cita lead time de 210 dias da SEL para IEDs/relés – validar com Suprimentos antes do PRM.').font = font(italic=True, color='C00000')
wa.cell(ni + 4, 1, 'Contratação da EPS: O.S. assinada antes da 1ª execução (premissa do template KOM).').font = font(italic=True, color='4D4D4D')
wa.freeze_panes = 'A4'

# ============================================================ Base PN 2026 (evidência)
wr = wb.create_sheet('Base PN 2026 (RL-RT)', 5)
wr['A1'] = 'Evidência dos prazos – blocos de instalação de RL e RT do WBS PN 2026 (arquivo WBS.xlsx)'; wr['A1'].font = font(bold=True, size=14, color=NAVY)
header(wr, 3, ['Linha WBS.xlsx', 'Bloco (PN 2026)', 'Equip.', 'Executar Obras (dias)', 'Duração total (dias)', 'Início', 'Término', 'Usado como referência?'], [12, 100, 8, 14, 14, 12, 12, 38], 36)
ref = [
 (54, 'BLOCO 1 – Bay e SDMT AL 10 SE Palmas III – Instalação de 5 RL e remanejamento de 1 RL 13,8 kV (71018)', 6, 45, 101, date(2026, 8, 10), date(2026, 11, 19), 'Sim – pacote 5+ equipamentos (P45)'),
 (105, 'BLOCO 3 – LDMT AL 2 de Araguaína I – Instalação de 5 RL 13,8 kV (73521) – Parte III', 5, 45, 101, date(2026, 6, 22), date(2026, 10, 1), 'Sim – pacote 5+ equipamentos (P45)'),
 (146, 'LD Bielândia – Palmeirante – Instalação de um regulador de tensão 34,5 kV (72240)', 1, 30, 86, date(2026, 7, 7), date(2026, 10, 1), 'Sim – pacote até 4 (P30)'),
 (157, 'LD Araguaína – Araguanã – Instalação de um regulador de tensão 34,5 kV (71579)', 1, 30, 86, date(2026, 7, 7), date(2026, 10, 1), 'Sim – pacote até 4 (P30)'),
 (188, 'AL SE Bielândia – Realocar RL (1), Instalar Religador 34,5 kV (2), Instalar Regulador de Tensão 34,5 kV (1) (71139) – Parte III', 4, 30, 87, date(2026, 6, 29), date(2026, 9, 24), 'Sim – pacote até 4 (P30)'),
 (219, 'LD Paranã – Palmeirópolis – Remanejamento de Regulador de Tensão 34,5 kV (72938)', 1, 30, 87, date(2026, 5, 11), date(2026, 8, 6), 'Sim – pacote até 4 (P30)'),
 (240, 'SDMT AL 01 de Palmeirópolis – Remanejar 1 Religador 13,8 kV (81319) – Parte II', 1, 30, 87, date(2026, 7, 13), date(2026, 10, 8), 'Sim – pacote até 4 (P30)'),
 (65, 'BLOCO 2 – Instalação de 4 Banco Capacitor 300 kVAr (71018)', 4, 30, 87, date(2026, 8, 10), date(2026, 11, 5), 'Não – banco de capacitores (fora do escopo RL/RT)'),
 (126, 'LDMT 1 e 6 de Araguaína III – Instalação de 6 bancos de capacitores (81179) – Parte II', 6, 30, 87, date(2026, 6, 15), date(2026, 9, 10), 'Não – banco de capacitores (fora do escopo RL/RT)'),
]
for i, r_ in enumerate(ref, 4):
    for j, v in enumerate(r_, 1):
        c = wr.cell(i, j, v); c.font = font(color='4D4D4D' if r_[7].startswith('Não') else '000000'); c.border = box
        c.alignment = Alignment(vertical='center', wrap_text=True, horizontal='left' if j in (2, 8) else 'center')
        if j in (6, 7): c.number_format = DF
    wr.row_dimensions[i].height = 28
nr = 3 + len(ref)
wr.cell(nr + 2, 1, 'Não utilizados: Planejamento (PMA, projeto executivo, site survey), obras de construção/recapacitação de rede (120–150 dias de execução) e bancos de capacitores.').font = font(italic=True, color='4D4D4D')
wr.cell(nr + 3, 1, 'Etapas comuns a todos os blocos RL/RT: Fiscalização 12 | Cadastro EO 10 | TCI/Integração final 15 | GAT 8 | Comissionar 15 | Enc. Técnico 7 (úteis) | Enc. Contábil 3 (úteis).').font = font(italic=True, color='4D4D4D')

wb.move_sheet('Parâmetros', offset=len(wb.sheetnames))
for _ws in wb.worksheets:
    _ws.page_setup.orientation = 'landscape'; _ws.page_setup.paperSize = 9
    _ws.page_setup.fitToWidth = 1; _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr.fitToPage = True
    _ws.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT)
print('saved', OUT, 'rows WBS', blk_last)
import json
json.dump({str(s): {k: v for k, v in R[s].items()} for s in R}, open('rows.json', 'w'))
