"""Ajustes no KOM do usuário: nome do projeto (topo e rodapé) + premissas/restrições/penalidades/riscos com base nas lições aprendidas."""
import sys, copy
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
NAME = 'Aquisição e Instalação de RL, RT e TripSaver'
NAME_UP = NAME.upper()
prs = Presentation(SRC)
S = list(prs.slides)

def shp(slide, name, nth=0, left=None):
    c = [s for s in slide.shapes if s.name == name and (left is None or s.left == left)]
    return c[nth]

def set_runs(par, text):
    rs = par.runs
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)

def set_paras(sh, texts, size=None):
    tb = sh.text_frame._txBody
    ps = tb.findall(qn('a:p'))
    while len(ps) < len(texts):
        n = copy.deepcopy(ps[-1]); ps[-1].addnext(n); ps = tb.findall(qn('a:p'))
    for i, p_ in enumerate(ps):
        if i < len(texts): set_runs(sh.text_frame.paragraphs[i], texts[i])
        else: tb.remove(p_)
    if size:
        for pa in sh.text_frame.paragraphs:
            for r in pa.runs: r.font.size = Pt(size)

# ---------------------------------------------------------------- 1) nome do projeto: topo (eyebrow) e rodapé
ref_eyebrow = shp(S[6], 'Eyebrow')            # slide 7: laranja + cinza
def fix_eyebrow(sl):
    eb = shp(sl, 'Eyebrow')
    if len(eb.text_frame.paragraphs[0].runs) < 2:   # eyebrow de run único -> copia o padrão (2 runs) do slide 7
        txb = eb.text_frame._txBody
        for p_ in txb.findall(qn('a:p')): txb.remove(p_)
        txb.append(copy.deepcopy(ref_eyebrow.text_frame._txBody.find(qn('a:p'))))
    rs = eb.text_frame.paragraphs[0].runs
    rs[0].text = 'KICKOFF MEETING'
    rs[1].text = f'   |   PN 2027  ·  {NAME_UP}'
    for r in rs[2:]: r._r.getparent().remove(r._r)
    eb.width = Emu(10800000)

for i in (2, 3, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14):          # slides 2,3,5..14 (índice 1-based)
    fix_eyebrow(S[i - 1])
# slide 15 (Obrigado)
eb = shp(S[14], 'Eyebrow'); set_runs(eb.text_frame.paragraphs[0], f'KICKOFF MEETING  ·  PN 2027  ·  {NAME_UP}'); eb.width = Emu(8200000)
for _r in eb.text_frame.paragraphs[0].runs: _r.font.size = Pt(11)
# slide 4 (resumo): título com o nome do projeto
set_runs(shp(S[3], 'Title 2').text_frame.paragraphs[0], f'PN 2027 – {NAME}')
# slide 10: nó raiz da EAP
tb = shp(S[9], 'TextBox 9')
set_runs(tb.text_frame.paragraphs[1], 'PN 2027')
set_runs(tb.text_frame.paragraphs[2], NAME)
# slide 9: nota (BC fora do escopo)
for sh in S[8].shapes:
    if sh.has_text_frame and 'RL/RT/BC' in sh.text_frame.text:
        for r in sh.text_frame.paragraphs[0].runs: r.text = r.text.replace('RL/RT/BC', 'RL/RT')
# rodapés: nome padronizado e numeração = posição real do slide
ref_footer = shp(S[6], 'Rodapé')
for idx, sl in enumerate(S, 1):
    has = [s for s in sl.shapes if s.name == 'Rodapé']
    if not has and idx == 6:                                   # curva orçamentária sem rodapé
        el = copy.deepcopy(ref_footer._element); sl.shapes._spTree.append(el)
        has = [s for s in sl.shapes if s.name == 'Rodapé']
    for f in has:
        set_runs(f.text_frame.paragraphs[0], f'{idx:02d}  ·  {NAME_UP}  ·  PN 2027')

# ---------------------------------------------------------------- 2) lições aprendidas -> premissas / restrições / penalidades / riscos
# Slide 7 — Premissas (11 posições)
prem = [
 'O projeto será executado conforme orçamento aprovado (base CAPEX 08/10/2026); desvios de rateio de materiais e impostos (referência: até 10%) são absorvidos no orçamento do departamento.',
 'EPS com capacidade executiva comprovada para o horizonte de 90 dias; mobilizará somente após geração e assinatura da O.S.',
 'Cumprimento dos prazos dos blocos de instalação de equipamentos especiais (RL/RT) definidos na WBS.',
 'Contratação de empresas com capacidade técnica e financeira, considerando as obras reguladas concorrentes na mesma empreiteira.',
 'Aprovação das LSS e análise da documentação de terceiros pelo SESMT ETO (10 dias corridos) dentro do prazo solicitado.',
 'Manutenção do escopo do projeto, sem inclusão de serviços não previstos inicialmente.',
 'Ausência de acidentes de trabalho que ocasionem paralisação; EPS atende a todos os procedimentos de segurança e às ITs.',
 'PMA dos equipamentos emitido logo após a aprovação do PN, com almoxarifado de destino conferido e kit completo de materiais periféricos.',
 'Local de instalação, autorização de passagem, medição de aterramento e sinal GPRS validados (ASPO, COPC e Operação) antes da execução.',
 'Equipamentos homologados na base de ativos e laudos validados antes do cadastro; DEOP, DMSL e GAT disponíveis para proteção e comissionamento.',
 'Disponibilidade de equipe de Linha Viva ETO e de desligamentos programados aprovados, garantindo continuidade nas áreas energizadas.']
for i, t in enumerate(prem, 1):
    sh = shp(S[6], f'Premissa {i}')
    set_runs(sh.text_frame.paragraphs[0], t)
    for r in sh.text_frame.paragraphs[0].runs: r.font.size = Pt(9.5)
restr = ['Iniciar a obra somente após assinatura do contrato de prestação de serviços com a empreiteira executora.',
         'Energização dentro do mês definido pela ASPO (abr, mai, jun e ago/2027).',
         'Instalar equipamentos somente com o local e o desligamento programado validados.',
         'Encerramento técnico antes do inventário anual do almoxarifado; contábil após fiscalização da Gestão de Ativos e até 17/11/2027.']
pen = ['Aplicação de multa às prestadoras de serviço e fornecedores, após notificação, em caso de descumprimento do contrato.',
       'Substituição da empreiteira e remanejamento de equipes por baixa capacidade ou atraso; medição pelas O.S. se faltar medição no prazo.',
       'Não atingimento do benefício esperado pela unidade.']
set_paras(shp(S[6], 'Restrições lista'), restr, 9)
set_paras(shp(S[6], 'Penalidades lista'), pen, 9.5)

# Slide 4 — resumo
set_paras(shp(S[3], 'Premissas lista'), [
 'O projeto será executado conforme orçamento aprovado (base CAPEX 08/10/2026); desvios de rateio/impostos (até 10%) absorvidos pelo departamento.',
 'EPS com capacidade executiva (90 dias) mobilizará somente após geração e assinatura da O.S.',
 'Execução de cada bloco em 30 dias (45 dias para o bloco com 4 RL).',
 'Equipamentos recebidos até 21/01, 19/03 e 28/05/2027 (lotes 1, 2 e 3).',
 'Aprovação das LSS e análise de documentos de terceiros (10 dias corridos) pelo SESMT ETO no prazo.',
 'Manutenção do escopo das fichas 2027, sem serviços não previstos.',
 'Ausência de acidentes com paralisação; EPS atende aos procedimentos de segurança e ITs.',
 'PMA emitido após a aprovação do PN, com almoxarifado de destino conferido e kit completo de periféricos.',
 'Local, autorização de passagem, aterramento e sinal GPRS validados antes da execução.',
 'Equipamentos homologados e laudos validados; DEOP, DMSL e GAT disponíveis para proteção e comissionamento.',
 'Disponibilidade de equipe de Linha Viva ETO e de desligamentos programados aprovados.'], 7.5)
set_paras(shp(S[3], 'Restrições lista'), [
 'Iniciar a obra somente após assinatura do contrato com a empreiteira executora.',
 'Energização dentro do mês definido pela ASPO (abr, mai, jun e ago/2027).',
 'Instalar somente com local e desligamento validados.',
 'Encerramento técnico antes do inventário; contábil após fiscalização da Gestão de Ativos, até 17/11/2027.'], 7.5)
set_paras(shp(S[3], 'Penalidades lista'), [
 'Multa às prestadoras de serviço e fornecedores, após notificação, em caso de descumprimento do contrato.',
 'Substituição da empreiteira e remanejamento de equipes por baixa capacidade ou atraso.',
 'Medição pelas O.S. se a empreiteira não enviar as medições no prazo.',
 'Não atingimento do benefício esperado pela unidade.'], 7.5)
set_paras(shp(S[3], 'Riscos lista'), [
 'Atraso no fornecimento de RL/RT e IEDs (lead time de até 210 dias – SEL).',
 'Prazo curto do 90811 (energização abr/27, execução a partir de 28/01/2027).',
 'Concentração de comissionamentos em jun/27 (6 blocos na Regional Centro).',
 'Indisponibilidade de Linha Viva / cancelamento de desligamentos programados.',
 'Local sem validação prévia ou empreiteira sem capacidade executiva.',
 'Acidente de trabalho com paralisação da obra.'], 7.5)

# Slide 8 — riscos
risks = [
 ('Acidente de trabalho', 'Acidente de trabalho durante a execução da obra',
  ['Reunião de integração com a empreiteira', 'Inspeções periódicas e DDS', 'Paradas de segurança trimestrais'],
  ['Paralisar qualquer atividade executada de forma insegura', 'Acionar o SESMT', 'Investigar o ocorrido']),
 ('Atraso em projeto, LSS e documentos', 'Projeto executivo, local sem autorização ou aterramento inadequado e LSS/documentos do SESMT',
  ['Validar o local com ASPO, COPC e Operação antes da execução', 'Mapear autorizações de passagem; medir o aterramento', 'Antecipar os documentos (SESMT: 10 dias)'],
  ['Remanejar o equipamento / definir novo ponto', 'Ampliar a malha de aterramento (poço/concreto)', 'Antecipar outras atividades da obra']),
 ('Atraso de materiais e equipamentos', 'RL/RT/TripSaver e kit de periféricos fora das datas da WBS (lead time de até 210 dias)',
  ['PMA logo após o PN; conferir o almoxarifado de destino', 'Diligenciar fornecedores e compras; kit completo', 'Acompanhar o forecast (rateio e impostos)'],
  ['Adiantar estruturas enquanto o equipamento não chega', 'Priorizar frentes pela meta de energização', 'Avaliar fornecedor homologado alternativo']),
 ('Linha Viva e desligamentos', 'Falta de equipe de Linha Viva ETO e cancelamento de desligamentos programados',
  ['Estudo prévio de viabilidade de desligamento e transferência de carga', 'Programar Linha Viva e COI com antecedência'],
  ['Intensificar o apoio de Linha Viva; reprogramar desligamentos', 'Ajustar a sequência das frentes']),
 ('Atraso no comissionamento', 'Laudos inconsistentes, equipamento não homologado, falha no EO/TS e comissionamentos concentrados em jun/27',
  ['Fluxo completo antes da SS ao DMSL (aterramento, EO/SIGOT, proteção, SCADA)', 'Validar laudos e homologação; reunião semanal'],
  ['Comissionar em bancada / linha de produção', 'Chamado à TI/GTD para cadastro; reavaliar a programação']),
 ('Atraso na execução / abandono de contrato', 'Baixa capacidade, desmobilização ou fim de contrato da empreiteira, com obras reguladas concorrentes',
  ['Acompanhar capacidade executiva (90 dias) e cronograma semanal', 'Evitar frentes em contratos próximos do fim; abertura de RFP'],
  ['Multa/advertência; remanejar equipes entre regionais', 'Aditivo de contratos vigentes ou substituição da empresa']),
]
for i, (tt, dd, mit, cont) in enumerate(risks, 1):
    t = shp(S[7], f'Risco {i} – título')
    set_runs(t.text_frame.paragraphs[0], tt); set_runs(t.text_frame.paragraphs[1], dd)
    pl = sorted([s for s in S[7].shapes if s.name == f'Risco {i} – plano'], key=lambda s: s.left)
    set_paras(pl[0], mit, 9); set_paras(pl[1], cont, 9)
    t.text_frame.paragraphs[1].runs[0].font.size = Pt(9)

# notas do orador: rastreabilidade das lições
notes7 = ('Base: Lições Aprendidas até 2025 + CM Obras de Distribuição Centro, Norte e Sul.\n'
 '- PMA/lead time/almoxarifado: ESE 2020 Instalação de Religadores; ESS 2021 Reguladores; ETO Obras Dist. Regional Sul/Centro 2020 (atraso de equipamentos especiais 34,5 kV).\n'
 '- Margem de 10% p/ rateio e impostos; kit de periféricos: ETO DEOP 2024-09 Aquisição e Instalação de Equipamentos Especiais.\n'
 '- Local/aterramento/passagem/GPRS: ETO DCMD 2024-01 e 2024-03; 2019 Sul; ESS 2021 Reguladores (DER, GPRS).\n'
 '- Homologação e laudos: ETO DEOP 2024-09 (Vacufuse); DCMD 2024-01/03 (laudos).\n'
 '- Capacidade da empreiteira, substituição, multa: ETO DCMD 2021-022, 2022-012/014, 2023-07/08; CM: encerramento de contrato de empreiteiras.\n'
 '- Medição pelas O.S.: EMT 2017 Instalação de Religadores.\n'
 '- Inventário e encerramento: ESE 2020; ETO DCMD 2019-009 (contábil até novembro).\n'
 '- Desligamentos: ETO DCMD 2022-012/013 (cancelamentos); CM: viabilidade de desligamento (atraso de 90 dias).')
S[6].notes_slide.notes_text_frame.text = notes7
S[7].notes_slide.notes_text_frame.text = ('Riscos que ocorreram no projeto anterior (CM): acidente de trabalho, atraso de projetos/documentos, atraso na execução, comissionamento, abandono de contrato. '
 'Novos riscos incorporados das lições: local/aterramento/autorização, cancelamento de desligamentos, laudos/homologação/EO-TS, lead time e rateio. Ver notas do slide 7.')
prs.save(OUT)
print('saved', OUT)
