import sys, math, datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
SRC, OUT = sys.argv[1:3]
# ---- datas do cronograma (WBS v3)
RB = openpyxl.load_workbook(SRC, data_only=True)['Resumo_Blocos']
blk = {RB.cell(r, 1).value: dict(ini=RB.cell(r, 6).value, op=RB.cell(r, 9).value) for r in range(4, 17)}

def utm_to_ll(E, N, zone, south=True):
    a, f = 6378137.0, 1 / 298.257222101; k0 = 0.9996
    e2 = f * (2 - f); ep2 = e2 / (1 - e2)
    x = E - 500000.0; y = N - (10000000.0 if south else 0.0)
    M = y / k0
    mu = M / (a * (1 - e2 / 4 - 3 * e2 ** 2 / 64 - 5 * e2 ** 3 / 256))
    e1 = (1 - math.sqrt(1 - e2)) / (1 + math.sqrt(1 - e2))
    p1 = (mu + (3 * e1 / 2 - 27 * e1 ** 3 / 32) * math.sin(2 * mu) + (21 * e1 ** 2 / 16 - 55 * e1 ** 4 / 32) * math.sin(4 * mu)
          + (151 * e1 ** 3 / 96) * math.sin(6 * mu) + (1097 * e1 ** 4 / 512) * math.sin(8 * mu))
    C1 = ep2 * math.cos(p1) ** 2; T1 = math.tan(p1) ** 2
    N1 = a / math.sqrt(1 - e2 * math.sin(p1) ** 2); R1 = a * (1 - e2) / (1 - e2 * math.sin(p1) ** 2) ** 1.5
    Dd = x / (N1 * k0)
    lat = p1 - (N1 * math.tan(p1) / R1) * (Dd ** 2 / 2 - (5 + 3 * T1 + 10 * C1 - 4 * C1 ** 2 - 9 * ep2) * Dd ** 4 / 24
                                           + (61 + 90 * T1 + 298 * C1 + 45 * T1 ** 2 - 252 * ep2 - 3 * C1 ** 2) * Dd ** 6 / 720)
    lon0 = math.radians(-183 + 6 * zone)
    lon = lon0 + (Dd - (1 + 2 * T1 + C1) * Dd ** 3 / 6 + (5 - 2 * C1 + 28 * T1 - 3 * C1 ** 2 + 8 * ep2 + 24 * T1 ** 2) * Dd ** 5 / 120) / math.cos(p1)
    return math.degrees(lat), math.degrees(lon)

# ---- relação (1 linha por ponto de instalação)
# (reg, polo, sigco, bloco, eq, kV, ação, projeto, cod_alim, se_fonte, topologia, tipo_ref, ref, utm(zona,E,N)|None, obs)
FICHA = 'ficha 2027'
P = [
 ('Centro', 'Paraíso', 70838, 'B8', 'RT', 34.5, 'Instalação', 'LD Cristalândia - Lagoa da Confusão', 'LD01025012', 'SE Cristalândia II', 'Anel', 'ID EO (poste)', '740431', None, 'Lado fonte: SE Cristalândia II.'),
 ('Centro', 'Paraíso', 70978, 'B10', 'RT', 34.5, 'Instalação', 'LD Paraíso I - Caseara', 'LD05062013', 'SE Paraíso I', 'Anel', 'Coordenada UTM', '720042.26 E / 8892930.53 S', (22, 720042.26, 8892930.53), 'Lado fonte: Paraíso do Tocantins; lado carga: Monte Santo. Fuso 22S assumido (a ficha não informa o fuso). Inclui balanceamento de fase (3 ramais) – sem equipamento de comunicação.'),
 ('Norte', 'Araguaína', 89249, 'B3', 'RT', 34.5, 'Instalação', 'LD Araguaína - Babaçulândia (2AR0)', 'não informado', 'SE Araguaína', 'não informada', 'ID EO', '535352', None, 'Lado fonte: SE Araguaína.'),
 ('Norte', 'Araguaína', 89267, 'B2', 'RT', 34.5, 'Instalação', 'LD Xambioá II - Araguanã (2XB2)', 'não informado', 'SE Xambioá II', 'não informada', 'ID EO', '298105', None, 'RT 1 de 2. Lado fonte: SE Xambioá II.'),
 ('Norte', 'Araguaína', 89267, 'B2', 'RT', 34.5, 'Instalação', 'LD Xambioá II - Araguanã (2XB2)', 'não informado', 'SE Xambioá II', 'não informada', 'ID EO', '38281041', None, 'RT 2 de 2. Lado fonte: SE Xambioá II.'),
 ('Centro', 'Palmas', 90811, 'B5', 'RT', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'ID EO (poste)', '33356216', None, 'RT 1 de 3.'),
 ('Centro', 'Palmas', 90811, 'B5', 'RT', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'ID EO (poste)', '498478', None, 'RT 2 de 3.'),
 ('Centro', 'Palmas', 90811, 'B5', 'RT', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'Chave', '8810033256', None, 'RT 3 de 3 (instalado no local da chave).'),
 ('Sul', 'Gurupi', 89307, 'B13', 'RT', 13.8, 'Instalação', 'AL 1 de Figueirópolis (2FGR)', 'AL01033047', 'SE Figueirópolis', 'não informada', 'ID EO (poste)', '47712358 (EO 241355)', None, 'Lado fonte: SE Figueirópolis; lado carga: chave 3100881036.'),
 ('Centro', 'Palmas', 90811, 'B4', 'RL', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'ID EO (poste)', '495670', None, 'RL 1 de 4.'),
 ('Centro', 'Palmas', 90811, 'B4', 'RL', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'Chave', '0312716256', None, 'RL 2 de 4. Instalado também para acionamento remoto do Reator BR12716256.'),
 ('Centro', 'Palmas', 90811, 'B4', 'RL', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'Chave', '8800010257', None, 'RL 3 de 4.'),
 ('Centro', 'Palmas', 90811, 'B4', 'RL', 34.5, 'Instalação', 'LD UHE Isamu - Ponte Alta (1UI1)', 'LD03046026', 'não informado', 'não informada', 'Chave', '3330781096', None, 'RL 4 de 4.'),
 ('Sul', 'Gurupi', 91979, 'B12', 'RL', 34.5, 'Instalação', 'LD Gurupi - Dueré (religador rural)', 'não informado', 'não informado', 'não informada', 'Chave', '3300246076', None, 'Religador rural. Os 3 TripSavers deste SIGCO constam na aba "TripSaver".'),
 ('Centro', 'Palmas', 92094, 'B11', 'RL', 34.5, 'Instalação', 'LD Palmas - Aparecida do Rio Negro', 'LD01059122', 'não informado', 'não informada', 'Chave', '8800756083', None, 'RL 1 de 3.'),
 ('Centro', 'Palmas', 92094, 'B11', 'RL', 34.5, 'Instalação', 'LD Palmas - Aparecida do Rio Negro', 'LD01059122', 'não informado', 'não informada', 'Chave', '8841714092', None, 'RL 2 de 3.'),
 ('Centro', 'Palmas', 92094, 'B11', 'RL', 34.5, 'Instalação', 'LD Palmas - Aparecida do Rio Negro', 'LD01059122', 'não informado', 'não informada', 'ID EO', '49973587', None, 'RL 3 de 3.'),
 ('Centro', 'Palmas', 92095, 'B6', 'RL', 34.5, 'Instalação', 'LD Palmas - Lajeado', 'LD03059122', 'não informado', 'não informada', 'ID EO (poste)', '425776', None, ''),
 ('Centro', 'Palmas', 92101, 'B7', 'RL', 13.8, 'Instalação', '8 de Palmas IV', 'AL08060122', 'SE Palmas IV (pelo nome do alimentador)', 'não informada', 'Chave', '8861680122', None, 'Religador instalado normalmente aberto (N.A).'),
 ('Centro', 'Palmas', 92102, 'B9', 'RL', 13.8, 'Instalação', '6 de Palmas IV', 'AL06060122', 'SE Palmas IV (pelo nome do alimentador)', 'não informada', 'Chave', '8803933122', None, 'Religador instalado normalmente fechado (N.F).'),
 ('Norte', 'Araguaína', 92118, 'B1', 'RL', 13.8, 'Instalação', '4 de Araguaína IV', 'não informado', 'SE Araguaína IV (pelo nome do alimentador)', 'não informada', 'ID EO', '551144', None, ''),
 ('Sul', 'Gurupi', 89307, 'B13', 'RL', 13.8, 'Remanejamento', 'AL 1 de Figueirópolis (2FGR)', 'AL01033047', 'SE Figueirópolis', 'não informada', 'Chave (destino)', '8857031047', None, 'Remanejamento do religador existente 7924620047 (automático N.F) para o ponto da chave 8857031047. Local atual do religador: não informado. Manobras: abrir CH 3300005047 e fechar CH 8857031047.'),
]
order = {'Norte': 0, 'Centro': 1, 'Sul': 2}
assert len(P) == 22
NAVY, TXT = '1F1F4C', '4D4D4D'
F = 'Arial'
fh = Font(name=F, bold=True, color='FFFFFF', size=10); fb = Font(name=F, size=10)
hdr = PatternFill('solid', fgColor=NAVY); thin = Side(style='thin', color='CCCCCC'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
alt = PatternFill('solid', fgColor='F3F5F8')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'RL_RT'
ws['A1'] = 'Relação de religadores (RL) e reguladores de tensão (RT) a instalar – PN 2027  |  Estudo de simulação de comunicação (Telecom)'
ws['A1'].font = Font(name=F, bold=True, size=13, color=NAVY)
ws['A2'] = ('Energisa Tocantins · DCMD · GP: Sandro Padilha Dias Junior · 12 SIGCOs · uma linha por ponto de instalação (13 RL e 9 RT, incluindo 1 RL remanejado). '
            'Fonte: fichas dos projetos 2027 e Cadastro CAPEX; datas: cronograma WBS (v3). Valores "não informado" não constam das fichas – ver aba Notas.')
ws['A2'].font = Font(name=F, italic=True, size=9, color=TXT); ws.merge_cells('A2:W2'); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 28
cols = ['Nº', 'Regional', 'Polo (validar com COCM)', 'SIGCO', 'Bloco (cronograma)', 'Equipamento', 'Tensão (kV)', 'Ação', 'Projeto / alimentador', 'Cód. alimentador',
        'SE fonte', 'Topologia', 'Tipo de referência', 'Ponto de referência (ID EO / chave / UTM)', 'UTM – fuso', 'UTM – E (m)', 'UTM – N (m)',
        'Latitude (convertida)', 'Longitude (convertida)', 'Início da execução prevista', 'Equip. em operação prevista', 'Observações']
widths = [5, 10, 14, 8, 11, 12, 9, 13, 36, 14, 26, 13, 18, 30, 8, 12, 13, 13, 13, 14, 14, 70]
HR = 4
for i, (t, w_) in enumerate(zip(cols, widths), 1):
    c = ws.cell(HR, i, t); c.font = fh; c.fill = hdr; c.border = box
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); ws.column_dimensions[L(i)].width = w_
ws.row_dimensions[HR].height = 42
rows = sorted(P, key=lambda p: (order[p[0]], p[1], p[2], p[3], p[4]))
for n, p in enumerate(rows, 1):
    reg, polo, sig, bl, eq, kv, acao, proj, cod, se, topo, tref, ref, utm, obs = p
    r = HR + n
    zone = E = N = lat = lon = None
    if utm:
        zone, E, N = utm; lat, lon = utm_to_ll(E, N, zone)
    vals = [n, reg, polo, sig, bl, eq, kv, acao, proj, cod, se, topo, tref, ref, (f'{zone}S' if zone else None), E, N,
            (round(lat, 6) if lat is not None else None), (round(lon, 6) if lon is not None else None), blk[bl]['ini'], blk[bl]['op'], obs]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = fb; c.border = box
        c.alignment = Alignment(vertical='center', wrap_text=i in (9, 11, 14, 22), horizontal='center' if i in (1, 2, 4, 5, 6, 7, 15, 16, 17, 18, 19, 20, 21) else 'left')
        if n % 2 == 0: c.fill = alt
    ws.cell(r, 7).number_format = '0.0'
    ws.cell(r, 16).number_format = '#,##0.00'; ws.cell(r, 17).number_format = '#,##0.00'
    ws.cell(r, 18).number_format = '0.000000'; ws.cell(r, 19).number_format = '0.000000'
    ws.cell(r, 20).number_format = 'dd/mm/yyyy'; ws.cell(r, 21).number_format = 'dd/mm/yyyy'
    ws.cell(r, 14).number_format = '@'
    ws.row_dimensions[r].height = 30 if len(obs) < 90 else 56
last = HR + len(rows)
ws.cell(last + 2, 1, 'Totais').font = Font(name=F, bold=True)
ws.cell(last + 3, 2, 'RL'); ws.cell(last + 3, 3, f'=COUNTIF(F{HR+1}:F{last},"RL")')
ws.cell(last + 4, 2, 'RT'); ws.cell(last + 4, 3, f'=COUNTIF(F{HR+1}:F{last},"RT")')
ws.cell(last + 5, 2, 'Total'); ws.cell(last + 5, 3, f'=C{last+3}+C{last+4}')
for rr in (last + 3, last + 4, last + 5):
    ws.cell(rr, 2).font = Font(name=F, bold=True); ws.cell(rr, 3).font = Font(name=F, bold=True)
ws.freeze_panes = ws.cell(HR + 1, 7)
ws.auto_filter.ref = f'A{HR}:{L(len(cols))}{last}'
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---- TripSaver (fora do estudo)
t = wb.create_sheet('TripSaver')
t['A1'] = 'TripSavers (religador monofásico) – SIGCO 91979 – informados apenas para completude (sem telecomando; fora do estudo de comunicação)'
t['A1'].font = Font(name=F, bold=True, size=12, color=NAVY)
for i, h in enumerate(['Nº', 'Regional', 'Polo (validar com COCM)', 'SIGCO', 'Bloco', 'Projeto / alimentador', 'Tipo de referência', 'Ponto de referência'], 1):
    c = t.cell(3, i, h); c.font = fh; c.fill = hdr; c.border = box; c.alignment = Alignment(horizontal='center', wrap_text=True, vertical='center')
for n, ch in enumerate(['0300839076', '3319901076', '3310406076'], 1):
    for i, v in enumerate([n, 'Sul', 'Gurupi', 91979, 'B12', 'LD Gurupi - Dueré', 'Chave', ch], 1):
        c = t.cell(3 + n, i, v); c.font = fb; c.border = box; c.alignment = Alignment(horizontal='center' if i != 6 else 'left'); 
        if i == 8: c.number_format = '@'
for i, w_ in enumerate([5, 10, 14, 8, 8, 30, 16, 18], 1): t.column_dimensions[L(i)].width = w_
t.sheet_view.showGridLines = False

# ---- Notas
nt = wb.create_sheet('Notas')
notes = [
 'Notas e premissas',
 '1. Cada linha da aba RL_RT é um ponto de instalação (um equipamento). Os 3 TripSavers do SIGCO 91979 estão na aba TripSaver (não possuem telecomando).',
 '2. Ponto de referência: ID EO (poste/elemento), código de chave ou coordenada UTM, conforme consta nas fichas dos projetos 2027. As fichas trazem coordenadas apenas para o RT do SIGCO 70978; para os demais pontos use o ID EO / chave no GIS para obter as coordenadas.',
 '3. Coordenadas UTM: SIGCO 70978 – a ficha não informa o fuso; adotado 22S (Paraíso do Tocantins) e datum SIRGAS 2000. Latitude/longitude calculadas por conversão UTM→geográfica (GRS80) e servem como referência; confirmar com a base GIS.',
 '4. "Polo" e "Regional" definidos pela localização dos alimentadores (Norte: Araguaína; Centro: Palmas e Paraíso; Sul: Gurupi) – validar com a coordenação COCM.',
 '5. "não informado / não informada": dado não consta nas fichas (cód. de alimentador, SE fonte ou topologia). A topologia (Anel) só está explícita nas fichas dos SIGCOs 70838 e 70978. Onde a SE fonte foi inferida do nome do alimentador, está indicado na própria célula.',
 '6. Datas: início da execução e operação prevista do bloco, conforme cronograma WBS (v3) – equipamentos disponíveis somente após 05/04/2027. Podem ser usadas como prazo para a entrega do estudo de comunicação.',
 '7. Remanejamento (SIGCO 89307): o religador existente 7924620047 (automático N.F, 13,8 kV) será levado para o ponto da chave 8857031047; o ponto listado é o destino.',
]
for i, tx in enumerate(notes, 1):
    c = nt.cell(i, 1, tx); c.alignment = Alignment(wrap_text=True, vertical='top'); c.font = Font(name=F, size=10, bold=(i == 1), color=NAVY if i == 1 else TXT)
nt.column_dimensions['A'].width = 140
for i in range(2, len(notes) + 1): nt.row_dimensions[i].height = 32
nt.sheet_view.showGridLines = False
wb.save(OUT)
# teste da conversão
print([ (round(a,5),round(b,5)) for a,b in [utm_to_ll(720042.26,8892930.53,22), utm_to_ll(342973,8836613,23)]])
